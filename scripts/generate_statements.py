"""Generate synthetic bank statements and a company ledger with planted discrepancies.

The real statements are confidential, so ReconPipe ships its own test data.
One run writes, for Q1 2026 (1 Jan - 31 Mar):

    data/raw/bank_a_statement_2026Q1.pdf   ruled table, DD/MM/YYYY, debit/credit columns
    data/raw/bank_b_statement_2026Q1.pdf   unruled columns, "15 Mar 2026", one signed amount
    data/raw/bank_c_statement_2026Q1.xlsx  noisy header rows, YYYY-MM-DD, money out/in columns
    data/raw/ledger.xlsx                   the company's books for all three accounts
    tests/expected_results.json            ground truth: every planted discrepancy, by id

The three bank layouts are deliberately different (column order, date format,
amount style), and the PDFs carry realistic noise: a logo header, page
footers, repeated column headers, brought/carried-forward lines, descriptions
that wrap onto a second table line, and one statement row printed twice
across a page break.

Planted in the ledger (rates are per bank transaction, per account):
    amount typo        ~2%  ledger amount differs (digits transposed / misplaced)
    missing in ledger  ~2%  bank line has no ledger entry
    missing in bank    ~1%  ledger entry never reached the bank (uncleared cheque)
    ledger duplicate   ~1%  the same entry posted twice
    date shift         ~2%  ledger date 4-9 days off and no reference: only the
                            fuzzy-description tier can match these
    one-to-many         1 per account: one bank transfer settles 2-3 ledger entries

Planted in the bank statements (data-quality defects the validate layer must
quarantine, one per business rule):
    Bank A  an extra "card authorisation reversal" line for 0.00   amount_non_zero
    Bank B  one line's balance misprinted                       balance_continuity
    Bank C  one line's year misprinted as 2027                  date_not_in_future
The B and C defects sit on lines that have no ledger entry anyway, so
quarantining them removes a missing-in-ledger exception rather than inventing
a new one. They are drawn after everything else, so planting them changes no
other transaction.

Everything is driven by one seeded random.Random, so the same seed always
produces the same transactions and the same expected_results.json.

Usage:
    python scripts/generate_statements.py [--seed 42] [--out-dir data/raw]
                                          [--expected tests/expected_results.json]
"""

from __future__ import annotations

import argparse
import io
import json
import random
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parent.parent

PERIOD_START = date(2026, 1, 1)
PERIOD_END = date(2026, 3, 31)
COMPANY = "Inzozi Trading Ltd"  # fictional account holder
DISCLAIMER = "Synthetic statement for software testing. This bank does not exist."

# Tier-2 matches on amount + date within this many days; date shifts are planted
# beyond it so that only the fuzzy-description tier can recover them.
TIER2_WINDOW_DAYS = 2
DATE_SHIFT_DAYS = (4, 9)

RATES = {
    "amount_typo": 0.02,
    "missing_in_ledger": 0.02,
    "missing_in_bank": 0.01,
    "ledger_duplicate": 0.01,
    "date_shift": 0.02,
}
ONE_TO_MANY_PER_ACCOUNT = 1
# Share of ordinary matched lines whose ledger entry carries the bank reference
# (tier 1). The rest have no reference and fall to amount + date (tier 2).
LEDGER_HAS_REF = 0.8


@dataclass(frozen=True)
class Bank:
    key: str
    name: str
    account_no: str
    layout: str
    file_name: str
    n_txns: int
    opening: Decimal


BANKS = [
    Bank(
        "A",
        "Kivu Commercial Bank",
        "0040-1187342-01",
        "pdf_ruled_debit_credit",
        "bank_a_statement_2026Q1.pdf",
        110,
        Decimal("18450000.00"),
    ),
    Bank(
        "B",
        "Akagera Trust Bank",
        "200-554-781-903",
        "pdf_unruled_signed_amount",
        "bank_b_statement_2026Q1.pdf",
        95,
        Decimal("7325500.00"),
    ),
    Bank(
        "C",
        "Nyungwe Cooperative Bank",
        "NCB-7781-0042",
        "xlsx_money_out_in",
        "bank_c_statement_2026Q1.xlsx",
        85,
        Decimal("3910250.00"),
    ),
]


@dataclass(frozen=True)
class Counterparty:
    name: str
    kind: str  # pos | transfer | debit_order | payroll | fee | receipt | interest
    purpose: str
    gl_account: str
    low: int
    high: int

    @property
    def outgoing(self) -> bool:
        return self.kind not in ("receipt", "interest")

    @property
    def named(self) -> bool:
        """Has a counterparty name in both bank and ledger text (fuzzy-matchable)."""
        return self.kind in ("pos", "transfer", "debit_order", "receipt")


COUNTERPARTIES = [
    Counterparty("Simba Supermarket", "pos", "office supplies", "6100", 8_000, 180_000),
    Counterparty("Muhazi Fuel Station", "pos", "vehicle fuel", "6300", 20_000, 250_000),
    Counterparty("Akagera Office Supplies Ltd", "transfer", "stationery", "6100", 60_000, 900_000),
    Counterparty(
        "Kigali Fresh Logistics Ltd", "transfer", "delivery services", "6200", 90_000, 1_400_000
    ),
    Counterparty(
        "Nyarutarama Property Management", "transfer", "office rent", "6400", 800_000, 2_500_000
    ),
    Counterparty("Ishema Power Utility", "debit_order", "electricity", "6500", 40_000, 320_000),
    Counterparty("Rugali Telecom", "debit_order", "internet and phone", "6510", 30_000, 180_000),
    Counterparty("Payroll", "payroll", "staff salaries", "6000", 3_000_000, 7_500_000),
    Counterparty("Bank charges", "fee", "account maintenance fee", "6900", 2_000, 15_000),
    Counterparty("Kivu Hotels Group Ltd", "receipt", "sales invoice", "4000", 250_000, 4_800_000),
    Counterparty("Gisenyi Beach Resort", "receipt", "sales invoice", "4000", 150_000, 2_600_000),
    Counterparty(
        "Huye Coffee Exporters Ltd", "receipt", "sales invoice", "4000", 400_000, 5_500_000
    ),
    Counterparty("Musanze Tours Ltd", "receipt", "sales invoice", "4000", 120_000, 1_900_000),
    Counterparty("Interest", "interest", "credit interest", "4900", 500, 9_000),
]
# How often each counterparty appears, in the same order as COUNTERPARTIES.
WEIGHTS = [14, 9, 8, 7, 3, 4, 4, 3, 3, 9, 8, 7, 7, 2]


@dataclass
class BankTxn:
    bank: Bank
    seq: int
    date: date
    cp: Counterparty
    amount: Decimal  # signed: money in positive
    ref: str | None
    description: str
    invoice: str | None
    balance: Decimal = Decimal("0")
    status: str = "matched"
    tier: str | None = None
    discrepancy: str | None = None
    ledger_ids: list[str] = field(default_factory=list)
    page: int | None = None
    excel_row: int | None = None
    wrapped: bool = False
    row_number: int | None = None  # what a parser emits: printed position or Excel row
    printed_date: date | None = None  # set only where the statement misprints the date
    printed_balance: Decimal | None = None  # set only where it misprints the balance
    quality: str | None = None  # the business rule a planted defect breaks
    extra: bool = False  # printed on the statement but not a real transaction


@dataclass
class LedgerEntry:
    entry_id: str
    date: date
    account_no: str
    ref: str | None
    description: str
    counterparty: str
    gl_account: str
    amount: Decimal
    status: str  # matched | amount_mismatch | missing_in_bank | duplicate_suspected
    bank_seq: int | None = None
    duplicate_of: str | None = None


# --------------------------------------------------------------------------- #
# Transaction generation
# --------------------------------------------------------------------------- #


def money(n: int) -> Decimal:
    return Decimal(n).quantize(Decimal("0.01"))


def fmt(amount: Decimal) -> str:
    """12500 -> '12,500.00' (absolute value)."""
    return f"{abs(amount):,.2f}"


def business_days(rng: random.Random, n: int) -> list[date]:
    days = []
    d = PERIOD_START
    while d <= PERIOD_END:
        if d.weekday() < 5:
            days.append(d)
        d += timedelta(days=1)
    return sorted(rng.choice(days) for _ in range(n))


def bank_reference(bank: Bank, rng: random.Random, d: date) -> str:
    if bank.key == "A":
        alnum = "".join(rng.choices("ABCDEFGHJKLMNPQRSTUVWXYZ23456789", k=5))
        return f"FT26{d.timetuple().tm_yday:03d}{alnum}"
    if bank.key == "B":
        return f"AKB{rng.randint(100000, 999999)}"
    return f"TX{rng.randint(10_000_000, 99_999_999)}"


def bank_description(
    bank: Bank, cp: Counterparty, rng: random.Random, invoice: str | None, d: date
) -> str:
    name = cp.name.upper()
    month = d.strftime("%b").upper()
    if bank.key == "A":
        return {
            "pos": f"POS {rng.randint(1000, 9999)} {name} KIGALI RW",
            "transfer": f"TRF TO {name} INV {invoice}",
            "debit_order": f"DD {name} ACC {rng.randint(10000, 99999)}",
            "payroll": f"SALARY BATCH {month} 2026",
            "fee": "MONTHLY ACCOUNT MAINTENANCE FEE",
            "receipt": f"TRF FROM {name} INV {invoice}",
            "interest": "CREDIT INTEREST",
        }[cp.kind]
    if bank.key == "B":
        return {
            "pos": f"Card purchase {cp.name} Kigali",
            "transfer": f"Outward transfer {cp.name} ref inv {invoice}",
            "debit_order": f"Direct debit {cp.name}",
            "payroll": f"Bulk salary payment {d.strftime('%B')} 2026",
            "fee": "Service charge",
            "receipt": f"Inward transfer from {cp.name} inv {invoice}",
            "interest": "Interest credited",
        }[cp.kind]
    return {
        "pos": f"Card: {cp.name}",
        "transfer": f"Payment: {cp.name} / INV-{invoice}",
        "debit_order": f"Debit order: {cp.name}",
        "payroll": f"Payroll run {month}-2026",
        "fee": "Ledger fee",
        "receipt": f"Receipt: {cp.name} / INV-{invoice}",
        "interest": "Interest",
    }[cp.kind]


def ledger_description(cp: Counterparty, invoice: str | None, bank: Bank) -> str:
    if cp.kind == "fee":
        return f"Bank charges - {bank.name}"
    if cp.kind == "interest":
        return f"Interest income - {bank.name}"
    if cp.kind == "payroll":
        return "Salaries - monthly payroll"
    if cp.kind == "receipt":
        return f"Receipt - {cp.name} - invoice {invoice}"
    suffix = f" - invoice {invoice}" if invoice else ""
    return f"Payment - {cp.name} - {cp.purpose}{suffix}"


def unique_amount(rng: random.Random, low: int, high: int, used: set[Decimal]) -> Decimal:
    """Amounts are unique per account so amount-based matching is never ambiguous."""
    while True:
        value = money(rng.randrange(low, high + 1, 50))
        if value not in used:
            used.add(value)
            return value


def typo_amount(amount: Decimal, rng: random.Random, used: set[Decimal]) -> Decimal:
    """A plausible keying error: two digits transposed, or one digit misplaced."""
    digits = str(int(abs(amount)))
    for _ in range(50):
        if len(digits) >= 3 and rng.random() < 0.7:
            i = rng.randrange(len(digits) - 1)
            if digits[i] == digits[i + 1]:
                continue
            wrong = digits[:i] + digits[i + 1] + digits[i] + digits[i + 2 :]
        else:
            i = rng.randrange(len(digits))
            wrong = digits[:i] + str((int(digits[i]) + rng.randint(1, 8)) % 10) + digits[i + 1 :]
        value = money(int(wrong))
        if value > 0 and value not in used:
            used.add(value)
            return value if amount > 0 else -value
    raise RuntimeError(f"could not build a typo for {amount}")


def generate_bank_txns(bank: Bank, rng: random.Random, invoice_counter: list[int]) -> list[BankTxn]:
    used: set[Decimal] = set()
    txns = []
    for d in business_days(rng, bank.n_txns):
        cp = rng.choices(COUNTERPARTIES, weights=WEIGHTS)[0]
        invoice = None
        if cp.kind in ("transfer", "receipt"):
            invoice_counter[0] += 1
            invoice = str(invoice_counter[0])
        value = unique_amount(rng, cp.low, cp.high, used)
        amount = -value if cp.outgoing else value
        ref = None if cp.kind in ("fee", "interest") else bank_reference(bank, rng, d)
        txns.append(
            BankTxn(
                bank, 0, d, cp, amount, ref, bank_description(bank, cp, rng, invoice, d), invoice
            )
        )
    balance = bank.opening
    for seq, t in enumerate(txns, start=1):
        t.seq = seq
        balance += t.amount
        t.balance = balance
    return txns


def pick(rng: random.Random, pool: list[BankTxn], k: int, taken: set[int]) -> list[BankTxn]:
    candidates = [t for t in pool if t.seq not in taken]
    chosen = rng.sample(candidates, min(k, len(candidates)))
    taken.update(t.seq for t in chosen)
    return chosen


def plant_and_build_ledger(
    bank: Bank, txns: list[BankTxn], rng: random.Random
) -> list[LedgerEntry]:
    """Decide every bank line's fate and write the ledger entries that follow from it."""
    used = {abs(t.amount) for t in txns}
    n = len(txns)
    count = {k: max(1, round(rate * n)) for k, rate in RATES.items()}
    taken: set[int] = set()

    one_to_many = pick(
        rng,
        [t for t in txns if t.cp.kind == "transfer" and t.ref and abs(t.amount) >= 200_000],
        ONE_TO_MANY_PER_ACCOUNT,
        taken,
    )
    shift_window = timedelta(days=DATE_SHIFT_DAYS[1] + 1)
    date_shift = pick(
        rng,
        [
            t
            for t in txns
            if t.cp.named
            and PERIOD_START + shift_window <= t.date
            and t.date <= PERIOD_END - shift_window
        ],
        count["date_shift"],
        taken,
    )
    typos = pick(
        rng, [t for t in txns if t.ref and abs(t.amount) >= 1000], count["amount_typo"], taken
    )
    missing = pick(rng, txns, count["missing_in_ledger"], taken)

    entries: list[LedgerEntry] = []

    def add(
        t: BankTxn | None,
        d: date,
        ref: str | None,
        amount: Decimal,
        status: str,
        description: str | None = None,
    ) -> LedgerEntry:
        cp = t.cp if t else None
        e = LedgerEntry(
            entry_id="",
            date=d,
            account_no=bank.account_no,
            ref=ref,
            description=description or ledger_description(cp, t.invoice, bank),
            counterparty=cp.name if cp else "",
            gl_account=cp.gl_account if cp else "",
            amount=amount,
            status=status,
            bank_seq=t.seq if t else None,
        )
        entries.append(e)
        if t:
            t.ledger_ids.append(id(e))  # replaced by the real entry_id after numbering
        return e

    for t in txns:
        lag = timedelta(days=rng.randint(0, TIER2_WINDOW_DAYS))
        if t in one_to_many:
            parts = split_amount(abs(t.amount), rng, used)
            for i, part in enumerate(parts, start=1):
                add(
                    t,
                    t.date - lag,
                    t.ref,
                    part if t.amount > 0 else -part,
                    "matched",
                    f"{ledger_description(t.cp, None, bank)} - invoice {t.invoice}-{i}",
                )
            t.tier, t.discrepancy = "tier1_exact", "one_to_many"
        elif t in date_shift:
            shift = timedelta(days=rng.randint(*DATE_SHIFT_DAYS))
            # Payments are booked before they clear; receipts are booked after they land.
            d = t.date - shift if t.amount < 0 else t.date + shift
            add(t, d, None, t.amount, "matched")
            t.tier, t.discrepancy = "tier3_fuzzy", "date_shift"
        elif t in typos:
            add(t, t.date - lag, t.ref, typo_amount(t.amount, rng, used), "amount_mismatch")
            t.status, t.discrepancy = "amount_mismatch", "amount_typo"
        elif t in missing:
            t.status, t.discrepancy = "missing_in_ledger", "missing_in_ledger"
        elif t.ref and rng.random() < LEDGER_HAS_REF:
            add(t, t.date - timedelta(days=rng.randint(0, 4)), t.ref, t.amount, "matched")
            t.tier = "tier1_exact"
        else:
            add(t, t.date - lag, None, t.amount, "matched")
            t.tier = "tier2_amount_date"

    # Ledger duplicates: an ordinary matched entry posted a second time.
    ordinary = [
        e
        for e in entries
        if e.status == "matched" and txn_by_seq(txns, e.bank_seq).discrepancy is None
    ]
    for original in rng.sample(ordinary, count["ledger_duplicate"]):
        dup = LedgerEntry(
            **{**original.__dict__, "status": "duplicate_suspected", "bank_seq": None}
        )
        dup.duplicate_of = id(original)
        entries.append(dup)

    # Missing in bank: cheques written late in the quarter that have not cleared.
    supplier = [c for c in COUNTERPARTIES if c.kind == "transfer"]
    for _ in range(count["missing_in_bank"]):
        cp = rng.choice(supplier)
        d = PERIOD_END - timedelta(days=rng.randint(1, 6))
        amount = -unique_amount(rng, cp.low, cp.high, used)
        cheque = f"CHQ{rng.randint(100100, 100999)}"
        entries.append(
            LedgerEntry(
                "",
                d,
                bank.account_no,
                cheque,
                f"Payment - {cp.name} - {cp.purpose} - cheque {cheque[3:]}",
                cp.name,
                cp.gl_account,
                amount,
                "missing_in_bank",
            )
        )
    return entries


def txn_by_seq(txns: list[BankTxn], seq: int | None) -> BankTxn:
    return next(t for t in txns if t.seq == seq)


def split_amount(total: Decimal, rng: random.Random, used: set[Decimal]) -> list[Decimal]:
    """Split a payment into 2-3 invoice amounts (multiples of 50, none reused)."""
    k = rng.choice([2, 3])
    units = int(total) // 50
    while True:
        cuts = sorted(rng.sample(range(1, units), k - 1))
        parts = [money((b - a) * 50) for a, b in zip([0, *cuts], [*cuts, units], strict=True)]
        if all(p >= 10_000 and p not in used for p in parts) and len(set(parts)) == k:
            used.update(parts)
            return parts


def number_ledger(entries: list[LedgerEntry], txns: list[BankTxn]) -> None:
    """Sort by date and assign JE ids; replace temporary object ids with real ones."""
    entries.sort(key=lambda e: (e.date, e.account_no, e.bank_seq or 0, e.status))
    real_id = {}
    for i, e in enumerate(entries, start=1):
        e.entry_id = f"JE-2026-{i:04d}"
        if e.status != "duplicate_suspected":
            real_id[id(e)] = e.entry_id
    for e in entries:
        if e.duplicate_of is not None:
            e.duplicate_of = real_id[e.duplicate_of]
    for t in txns:
        t.ledger_ids = [real_id[x] for x in t.ledger_ids]


REVERSAL = Counterparty("Card authorisation reversal", "reversal", "", "", 0, 0)


def plant_quality_defects(
    all_txns: dict[str, list[BankTxn]], rng: random.Random
) -> dict[str, list[BankTxn]]:
    """Plant one statement defect per bank; return each bank's lines in printed order."""
    printed = {key: list(txns) for key, txns in all_txns.items()}

    a = all_txns["A"]
    k = rng.randrange(5, len(a) - 5)
    reversal = BankTxn(
        bank=a[k].bank,
        seq=0,
        date=a[k].date,
        cp=REVERSAL,
        amount=Decimal("0.00"),
        ref=None,
        description="CARD AUTHORISATION REVERSAL",
        invoice=None,
        balance=a[k].balance,
        status="quarantined",
        quality="amount_non_zero",
        extra=True,
    )
    printed["A"].insert(k + 1, reversal)

    def unmatched(txns: list[BankTxn], avoid: set[int]) -> BankTxn:
        pool = [t for t in txns if t.status == "missing_in_ledger" and t.seq not in avoid]
        if not pool:
            raise RuntimeError(f"no missing-in-ledger line to plant a defect on in {txns[0].bank}")
        return rng.choice(pool)

    b = all_txns["B"]
    page_break = paginate(b, B_WRAP, B_LINES_PER_PAGE)[0][-1].seq  # reprinted line
    t = unmatched(b, {1, page_break, len(b)})
    t.printed_balance = t.balance + Decimal("9000.00")
    t.status, t.quality = "quarantined", "balance_continuity"

    t = unmatched(all_txns["C"], set())
    t.printed_date = t.date.replace(year=PERIOD_END.year + 1)
    t.status, t.quality = "quarantined", "date_not_in_future"
    return printed


# --------------------------------------------------------------------------- #
# Rendering: Bank A, ruled PDF table with debit and credit columns
# --------------------------------------------------------------------------- #

A_COLS = [
    ("Date", 36, 88),
    ("Description", 88, 300),
    ("Reference", 300, 380),
    ("Debit", 380, 445),
    ("Credit", 445, 505),
    ("Balance", 505, 568),
]
A_WRAP = 40  # characters per description line
A_LINE_H = 13
A_LINES_PER_PAGE = 46


def wrap(text: str, width: int) -> list[str]:
    if len(text) <= width:
        return [text]
    cut = text.rfind(" ", 0, width)
    cut = cut if cut > 0 else width
    return [text[:cut], text[cut:].strip()]


def paginate(txns: list[BankTxn], width: int, lines_per_page: int) -> list[list[BankTxn]]:
    pages: list[list[BankTxn]] = [[]]
    used = 0
    for t in txns:
        lines = len(wrap(t.description, width))
        if used + lines > lines_per_page:
            pages.append([])
            used = 0
        pages[-1].append(t)
        used += lines
    return pages


def render_bank_a(bank: Bank, txns: list[BankTxn], path: Path) -> None:
    c = Canvas(str(path), pagesize=A4, invariant=True)
    c.setTitle(f"{bank.name} statement {bank.account_no}")
    width, height = A4
    pages = paginate(txns, A_WRAP, A_LINES_PER_PAGE)
    navy = HexColor("#16365C")
    for page_no, page in enumerate(pages, start=1):
        # Logo header band
        c.setFillColor(navy)
        c.rect(0, height - 70, width, 70, stroke=0, fill=1)
        c.setFillColor(white)
        c.circle(52, height - 35, 18, stroke=0, fill=1)
        c.setFillColor(navy)
        c.setFont("Helvetica-Bold", 14)
        c.drawCentredString(52, height - 40, "KCB")
        c.setFillColor(white)
        c.setFont("Helvetica-Bold", 16)
        c.drawString(80, height - 32, bank.name.upper())
        c.setFont("Helvetica", 8)
        c.drawString(80, height - 46, "Avenue de la Paix, Kigali, Rwanda  |  Customer care 4040")
        c.setFillColor(HexColor("#000000"))

        y = height - 92
        c.setFont("Helvetica-Bold", 11)
        c.drawString(36, y, "STATEMENT OF ACCOUNT")
        c.setFont("Helvetica", 8)
        info = [
            f"Account holder: {COMPANY}",
            f"Account number: {bank.account_no}",
            "Currency: RWF",
            f"Period: {PERIOD_START:%d/%m/%Y} - {PERIOD_END:%d/%m/%Y}",
        ]
        for i, line in enumerate(info):
            c.drawString(36 + (i % 2) * 260, y - 14 - (i // 2) * 11, line)
        if page_no == 1:
            c.drawString(36, y - 40, f"Opening balance: {fmt(bank.opening)}")

        top = y - 52
        lines: list[tuple[list[str], bool]] = [
            (["Date", "Description", "Reference", "Debit", "Credit", "Balance"], True)
        ]
        for t in page:
            t.page = page_no
            first, *rest = wrap(t.description, A_WRAP)
            t.wrapped = bool(rest)
            debit = fmt(t.amount) if t.amount <= 0 else ""  # a 0.00 line prints as a debit
            credit = fmt(t.amount) if t.amount > 0 else ""
            lines.append(
                (
                    [
                        t.date.strftime("%d/%m/%Y"),
                        first,
                        t.ref or "",
                        debit,
                        credit,
                        fmt(t.balance),
                    ],
                    False,
                )
            )
            for extra in rest:  # wrapped description: its own ruled table line
                lines.append((["", extra, "", "", "", ""], False))

        bottom = top - A_LINE_H * len(lines)
        c.setLineWidth(0.4)
        for i in range(len(lines) + 1):
            c.line(A_COLS[0][1], top - i * A_LINE_H, A_COLS[-1][2], top - i * A_LINE_H)
        for _, x0, _ in A_COLS:
            c.line(x0, top, x0, bottom)
        c.line(A_COLS[-1][2], top, A_COLS[-1][2], bottom)
        for i, (cells, is_header) in enumerate(lines):
            c.setFont("Helvetica-Bold" if is_header else "Helvetica", 7.5)
            baseline = top - i * A_LINE_H - 9
            for (name, x0, x1), text in zip(A_COLS, cells, strict=True):
                if name in ("Debit", "Credit", "Balance") and not is_header:
                    c.drawRightString(x1 - 3, baseline, text)
                else:
                    c.drawString(x0 + 3, baseline, text)

        if page_no == len(pages):
            c.setFont("Helvetica-Bold", 8)
            c.drawString(36, bottom - 16, f"Closing balance: {fmt(txns[-1].balance)}")
        c.setFont("Helvetica", 7)
        c.drawString(36, 28, DISCLAIMER)
        c.drawRightString(width - 36, 28, f"Page {page_no} of {len(pages)}")
        c.showPage()
    c.save()


# --------------------------------------------------------------------------- #
# Rendering: Bank B, unruled PDF columns with one signed amount
# --------------------------------------------------------------------------- #

B_X = {"date": 40, "ref": 110, "details": 185, "amount_r": 470, "balance_r": 560}
B_WRAP = 52
B_LINE_H = 12
B_LINES_PER_PAGE = 50


def signed(amount: Decimal) -> str:
    return f"RWF -{fmt(amount)}" if amount < 0 else f"RWF {fmt(amount)}"


def render_bank_b(bank: Bank, txns: list[BankTxn], path: Path) -> list[int]:
    """Returns the seq of statement lines printed twice (page-break overlap)."""
    c = Canvas(str(path), pagesize=A4, invariant=True)
    c.setTitle(f"{bank.name} statement {bank.account_no}")
    width, height = A4
    pages = paginate(txns, B_WRAP, B_LINES_PER_PAGE)
    green = HexColor("#1E5B3A")
    duplicated: list[int] = []
    prev_balance = bank.opening
    for page_no, page in enumerate(pages, start=1):
        c.setFillColor(green)
        c.roundRect(40, height - 78, 46, 46, 8, stroke=0, fill=1)
        c.setFillColor(white)
        c.setFont("Helvetica-Bold", 18)
        c.drawCentredString(63, height - 61, "AT")
        c.setFillColor(green)
        c.setFont("Helvetica-Bold", 15)
        c.drawString(96, height - 50, bank.name)
        c.setFont("Helvetica-Oblique", 8)
        c.drawString(96, height - 62, "Growing together since 1998")
        c.setFillColor(HexColor("#000000"))
        c.setFont("Helvetica", 8)
        for i, line in enumerate(
            [
                f"Customer: {COMPANY}",
                f"A/C No. {bank.account_no}",
                f"From {PERIOD_START:%d %b %Y} to {PERIOD_END:%d %b %Y}",
            ]
        ):
            c.drawRightString(width - 40, height - 44 - i * 11, line)

        y = height - 105
        c.setFont("Helvetica-Bold", 8)
        c.drawString(B_X["date"], y, "Value Date")
        c.drawString(B_X["ref"], y, "Ref No.")
        c.drawString(B_X["details"], y, "Transaction Details")
        c.drawRightString(B_X["amount_r"], y, "Amount")
        c.drawRightString(B_X["balance_r"], y, "Balance")
        c.setLineWidth(0.6)
        c.line(40, y - 4, width - 35, y - 4)
        y -= B_LINE_H + 4

        c.setFont("Helvetica-Oblique", 8)
        c.drawString(B_X["details"], y, "BALANCE BROUGHT FORWARD")
        c.drawRightString(B_X["balance_r"], y, f"{fmt(prev_balance)} CR")
        y -= B_LINE_H

        rows = list(page)
        if page_no == 2:  # the last line of page 1 is printed again at the top of page 2
            rows = [pages[0][-1], *rows]
            duplicated.append(pages[0][-1].seq)
        for i, t in enumerate(rows):
            if not (page_no == 2 and i == 0):
                t.page = page_no
            first, *rest = wrap(t.description, B_WRAP)
            t.wrapped = bool(rest)
            c.setFont("Helvetica", 8)
            c.drawString(B_X["date"], y, t.date.strftime("%d %b %Y"))
            c.drawString(B_X["ref"], y, t.ref or "")
            c.drawString(B_X["details"], y, first)
            c.drawRightString(B_X["amount_r"], y, signed(t.amount))
            shown = t.balance if t.printed_balance is None else t.printed_balance
            c.drawRightString(B_X["balance_r"], y, f"{fmt(shown)} CR")
            y -= B_LINE_H
            for extra in rest:
                c.drawString(B_X["details"], y, extra)
                y -= B_LINE_H
        prev_balance = page[-1].balance

        c.setFont("Helvetica-Oblique", 8)
        label = "CLOSING BALANCE" if page_no == len(pages) else "BALANCE CARRIED FORWARD"
        c.drawString(B_X["details"], y, label)
        c.drawRightString(B_X["balance_r"], y, f"{fmt(prev_balance)} CR")
        c.setFont("Helvetica", 7)
        c.line(40, 42, width - 35, 42)
        c.drawString(40, 30, DISCLAIMER)
        c.drawRightString(width - 35, 30, f"{page_no}/{len(pages)}")
        c.showPage()
    c.save()
    return duplicated


# --------------------------------------------------------------------------- #
# Rendering: Bank C (Excel) and the ledger
# --------------------------------------------------------------------------- #

FIXED_TIMESTAMP = datetime(2026, 4, 1, 8, 0, 0)


def save_workbook(wb: Workbook, path: Path) -> None:
    wb.properties.creator = "ReconPipe generator"
    wb.properties.created = FIXED_TIMESTAMP
    buffer = io.BytesIO()
    wb.save(buffer)
    # An .xlsx is a zip, and zip stamps each entry with the current time. Re-pack
    # with a fixed time so the same seed always gives byte-identical files.
    with (
        zipfile.ZipFile(buffer) as src,
        zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, FIXED_TIMESTAMP.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            data = src.read(item.filename)
            if item.filename == "docProps/core.xml":  # openpyxl stamps "modified" at save
                stamp = FIXED_TIMESTAMP.isoformat() + "Z"
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>" + stamp.encode(), data)
            dst.writestr(info, data)


def render_bank_c(bank: Bank, txns: list[BankTxn], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append([bank.name.upper()])
    ws.append(["Account statement"])
    ws.append([f"Account number: {bank.account_no}", None, f"Holder: {COMPANY}"])
    ws.append([f"Period: {PERIOD_START.isoformat()} to {PERIOD_END.isoformat()}"])
    ws.append([])
    ws.append(["Posting Date", "Narrative", "Txn Ref", "Money Out", "Money In", "Balance"])
    for cell in ws[1] + ws[6]:
        cell.font = Font(bold=True)
    ws.append([None, "Opening balance", None, None, None, float(bank.opening)])
    for t in txns:
        ws.append(
            [
                (t.printed_date or t.date).isoformat(),
                t.description,
                t.ref,
                abs(t.amount) if t.amount < 0 else None,
                t.amount if t.amount > 0 else None,
                t.balance,
            ]
        )
        t.excel_row = ws.max_row
    ws.append([])
    ws.append([None, "Closing balance", None, None, None, txns[-1].balance])
    ws.append(["End of statement"])
    for col, w in zip("ABCDEF", [13, 46, 13, 13, 13, 15], strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=7, min_col=4, max_col=6):
        for cell in row:
            cell.number_format = "#,##0.00"
    save_workbook(wb, path)


def render_ledger(entries: list[LedgerEntry], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(
        [
            "entry_id",
            "posting_date",
            "bank_account",
            "reference",
            "description",
            "counterparty",
            "gl_account",
            "amount",
        ]
    )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for e in entries:
        ws.append(
            [
                e.entry_id,
                e.date,
                e.account_no,
                e.ref,
                e.description,
                e.counterparty,
                e.gl_account,
                e.amount,
            ]
        )
        ws.cell(row=ws.max_row, column=2).number_format = "yyyy-mm-dd"
        ws.cell(row=ws.max_row, column=8).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGH", [14, 13, 17, 16, 58, 30, 11, 15], strict=True):
        ws.column_dimensions[col].width = w
    save_workbook(wb, path)


# --------------------------------------------------------------------------- #
# Ground truth
# --------------------------------------------------------------------------- #


def build_expected(
    seed: int,
    all_txns: dict[str, list[BankTxn]],
    entries: list[LedgerEntry],
    bank_dupes: dict[str, list[int]],
    printed: dict[str, list[BankTxn]],
    out_dir: Path,
) -> dict:
    txns = [t for ts in all_txns.values() for t in ts]
    defects = [t for ts in printed.values() for t in ts if t.quality]
    statuses = [t.status for t in txns] + [e.status for e in entries if e.bank_seq is None]

    def n(pred) -> int:
        return sum(1 for x in txns if pred(x))

    def rel(p: Path) -> str:
        try:
            return p.relative_to(ROOT).as_posix()
        except ValueError:
            return p.as_posix()

    return {
        "seed": seed,
        "period": {"start": PERIOD_START.isoformat(), "end": PERIOD_END.isoformat()},
        "currency": "RWF",
        "sign_convention": "credit (money in) positive, debit (money out) negative",
        "tier2_window_days": TIER2_WINDOW_DAYS,
        "statements": [
            {
                "file": rel(out_dir / b.file_name),
                "bank": b.name,
                "account_id": b.account_no,
                "layout": b.layout,
                "transactions": len(all_txns[b.key]),
                "printed_rows": len(printed_seqs(printed[b.key], bank_dupes.get(b.key, []))),
                "printed_seqs": printed_seqs(printed[b.key], bank_dupes.get(b.key, [])),
                "duplicate_printed_seqs": bank_dupes.get(b.key, []),
                "wrapped_descriptions": sum(t.wrapped for t in printed[b.key]),
                "pages": max((t.page or 0) for t in printed[b.key]) or None,
                "opening_balance": str(b.opening),
                "closing_balance": str(all_txns[b.key][-1].balance),
            }
            for b in BANKS
        ],
        "ledger": {"file": rel(out_dir / "ledger.xlsx"), "entries": len(entries)},
        "totals": {
            "bank_transactions": len(txns),
            "ledger_entries": len(entries),
        },
        "planted": {
            "amount_typo": n(lambda t: t.discrepancy == "amount_typo"),
            "missing_in_ledger": n(lambda t: t.discrepancy == "missing_in_ledger"),
            "missing_in_bank": sum(e.status == "missing_in_bank" for e in entries),
            "ledger_duplicate": sum(e.status == "duplicate_suspected" for e in entries),
            "date_shift": n(lambda t: t.discrepancy == "date_shift"),
            "one_to_many": n(lambda t: t.discrepancy == "one_to_many"),
            "bank_duplicate_printed_rows": sum(len(v) for v in bank_dupes.values()),
            "statement_quality_defects": len(defects),
        },
        "expected_quarantine": {
            "rows": len(defects),
            "by_rule": {
                r: sum(t.quality == r for t in defects)
                for r in sorted({t.quality for t in defects})
            },
        },
        "expected_matches": {
            "tier1_exact": n(lambda t: t.tier == "tier1_exact"),
            "tier2_amount_date": n(lambda t: t.tier == "tier2_amount_date"),
            "tier3_fuzzy": n(lambda t: t.tier == "tier3_fuzzy"),
            "total_matched_bank_lines": n(lambda t: t.status == "matched"),
        },
        "expected_classification": {s: statuses.count(s) for s in sorted(set(statuses))},
        "bank_transactions": [
            {
                "file": t.bank.file_name,
                "account_id": t.bank.account_no,
                "seq": t.seq,
                "row_number": t.row_number,
                "page": t.page,
                "excel_row": t.excel_row,
                "date": t.date.isoformat(),
                "printed_date": t.printed_date.isoformat() if t.printed_date else None,
                "printed_balance": None if t.printed_balance is None else str(t.printed_balance),
                "data_quality": t.quality,
                "ref_id": t.ref,
                "description": t.description,
                "amount": str(t.amount),
                "balance": str(t.balance),
                "expected_status": t.status,
                "expected_tier": t.tier,
                "discrepancy": t.discrepancy,
                "ledger_entry_ids": t.ledger_ids,
            }
            for t in txns
        ],
        "ledger_only": [
            {
                "entry_id": e.entry_id,
                "account_id": e.account_no,
                "date": e.date.isoformat(),
                "ref_id": e.ref,
                "amount": str(e.amount),
                "expected_status": e.status,
                "duplicate_of": e.duplicate_of,
            }
            for e in entries
            if e.bank_seq is None
        ],
        "quality_defects": [
            {
                "file": t.bank.file_name,
                "account_id": t.bank.account_no,
                "row_number": t.row_number,
                "seq": None if t.extra else t.seq,
                "rule": t.quality,
                "kind": "extra_zero_amount_line"
                if t.extra
                else ("misprinted_balance" if t.printed_balance is not None else "misprinted_date"),
                "date": (t.printed_date or t.date).isoformat(),
                "description": t.description,
                "amount": str(t.amount),
                "balance": str(t.balance if t.printed_balance is None else t.printed_balance),
            }
            for t in defects
        ],
    }


def printed_seqs(lines: list[BankTxn], duplicated: list[int]) -> list[int | None]:
    """The seq of every printed line in order (None for a line that is not a transaction)."""
    out: list[int | None] = []
    for t in lines:
        out.append(None if t.extra else t.seq)
        if not t.extra and t.seq in duplicated:
            out.append(t.seq)
    return out


def assign_row_numbers(printed: dict[str, list[BankTxn]], bank_dupes: dict[str, list[int]]) -> None:
    """Set what each parser will emit as row_number: printed position, or the Excel row."""
    for key, lines in printed.items():
        if key == "C":
            for t in lines:
                t.row_number = t.excel_row
            continue
        position = 0
        for t in lines:
            position += 1
            t.row_number = position
            if not t.extra and t.seq in bank_dupes.get(key, []):
                position += 1  # the reprint takes the next printed position


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #


def generate(seed: int, out_dir: Path, expected_path: Path) -> dict:
    rng = random.Random(seed)
    out_dir.mkdir(parents=True, exist_ok=True)
    expected_path.parent.mkdir(parents=True, exist_ok=True)

    invoice_counter = [1040]
    all_txns: dict[str, list[BankTxn]] = {}
    entries: list[LedgerEntry] = []
    for bank in BANKS:
        txns = generate_bank_txns(bank, rng, invoice_counter)
        entries += plant_and_build_ledger(bank, txns, rng)
        all_txns[bank.key] = txns
    number_ledger(entries, [t for ts in all_txns.values() for t in ts])
    printed = plant_quality_defects(all_txns, rng)

    bank_dupes: dict[str, list[int]] = {}
    render_bank_a(BANKS[0], printed["A"], out_dir / BANKS[0].file_name)
    bank_dupes["B"] = render_bank_b(BANKS[1], printed["B"], out_dir / BANKS[1].file_name)
    render_bank_c(BANKS[2], printed["C"], out_dir / BANKS[2].file_name)
    render_ledger(entries, out_dir / "ledger.xlsx")
    assign_row_numbers(printed, bank_dupes)

    expected = build_expected(seed, all_txns, entries, bank_dupes, printed, out_dir)
    expected_path.write_text(json.dumps(expected, indent=2) + "\n", encoding="utf-8")
    return expected


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--out-dir", type=Path, default=ROOT / "data" / "raw")
    parser.add_argument("--expected", type=Path, default=ROOT / "tests" / "expected_results.json")
    args = parser.parse_args(argv)

    expected = generate(args.seed, args.out_dir, args.expected)
    print(f"Wrote {len(BANKS)} statements + ledger to {args.out_dir}")
    for s in expected["statements"]:
        print(
            f"  {s['file']}: {s['transactions']} transactions, "
            f"{s['printed_rows']} printed rows, {s['wrapped_descriptions']} wrapped"
        )
    print(f"  ledger: {expected['ledger']['entries']} entries")
    print("Planted discrepancies:")
    for k, v in expected["planted"].items():
        print(f"  {k:<28} {v}")
    print("Statement defects to quarantine:")
    for k, v in expected["expected_quarantine"]["by_rule"].items():
        print(f"  {k:<28} {v}")
    print("Expected matches:")
    for k, v in expected["expected_matches"].items():
        print(f"  {k:<28} {v}")
    print(f"Ground truth: {args.expected}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
