"""Read the company ledger export into LedgerEntry rows.

The ledger is not a bank statement, so it is not in the parser registry (which
correctly skips it). It is an accounting-system export with one header row:

    entry_id | posting_date | bank_account | reference | description |
    counterparty | gl_account | amount

Columns are found by header name. Amounts go through Decimal(str(...)) so a
float cell never leaks binary rounding into money.
"""

from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from reconpipe.extract.base import ParseError
from reconpipe.extract.excel_parser import UNREADABLE, cell_text
from reconpipe.schemas import LedgerEntry

COLUMNS = {
    "entry_id": "entry_id",
    "posting_date": "date",
    "bank_account": "account_id",
    "reference": "ref_id",
    "description": "description",
    "counterparty": "counterparty",
    "gl_account": "gl_account",
    "amount": "amount",
}


def _date(value: object) -> dt.date | str | None:
    if isinstance(value, dt.datetime):
        return value.date()
    return value  # a date, an ISO string, or None: the model decides


def _money(value: object) -> Decimal | str | None:
    if isinstance(value, int | float):
        return Decimal(str(value))
    return cell_text(value)


def read_ledger(path: Path) -> list[LedgerEntry]:
    path = Path(path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except UNREADABLE as exc:
        raise ParseError(path, f"cannot open ledger: {exc}") from exc
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise ParseError(path, "ledger is empty")

    header = [str(c).strip().casefold() if c is not None else "" for c in rows[0]]
    missing = set(COLUMNS) - set(header)
    if missing:
        raise ParseError(path, f"ledger is missing columns: {', '.join(sorted(missing))}")
    index = {COLUMNS[name]: header.index(name) for name in COLUMNS}

    entries = []
    for i, row in enumerate(rows[1:], start=2):
        if all(v is None for v in row):
            continue
        values = {field: row[j] if j < len(row) else None for field, j in index.items()}
        try:
            entries.append(
                LedgerEntry(
                    entry_id=cell_text(values["entry_id"]),
                    date=_date(values["date"]),
                    account_id=cell_text(values["account_id"]),
                    ref_id=cell_text(values["ref_id"]),
                    description=cell_text(values["description"]),
                    counterparty=cell_text(values["counterparty"]),
                    gl_account=cell_text(values["gl_account"]),
                    amount=_money(values["amount"]),
                    source_file=path.name,
                    row_number=i,
                )
            )
        except ValidationError as exc:
            raise ParseError(path, f"row {i}: {exc.errors()[0]['msg']}") from exc
    return entries
