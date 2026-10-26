"""Excel statement parser: a header row somewhere near the top, one transaction per row.

Built for the Nyungwe Cooperative Bank export (Bank C), which looks like:

    row 1-4   bank name, title, "Account number: NCB-7781-0042", period
    row 6     Posting Date | Narrative | Txn Ref | Money Out | Money In | Balance
    row 7     opening balance line (no date)
    row 8..   transactions
    footer    blank row, closing balance line, "End of statement"

Columns are found by header name, not position, so a reordered export still
parses. A row is a transaction if it has a date cell and a money cell; that
rule skips the opening/closing balance lines and footer text without having to
list them.
"""

from __future__ import annotations

import datetime as dt
import re
from decimal import Decimal
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from reconpipe.extract.base import BaseStatementParser, ParseError
from reconpipe.schemas import RawTransaction

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
# What openpyxl raises for a missing, corrupt or non-Excel file.
UNREADABLE = (OSError, BadZipFile, InvalidFileException, KeyError, ValueError)
HEADER_SCAN_ROWS = 15

# Normalised header text -> RawTransaction field.
HEADER_ALIASES = {
    "posting date": "date",
    "narrative": "description",
    "txn ref": "ref_id",
    "money out": "debit",
    "money in": "credit",
    "balance": "balance",
}
REQUIRED_FIELDS = {"date", "description", "debit", "credit"}
ACCOUNT_PATTERN = re.compile(r"account\s*(?:number|no\.?)\s*:?\s*(\S+)", re.IGNORECASE)


def _normalise(value: object) -> str:
    return " ".join(str(value).split()).casefold() if value is not None else ""


def cell_text(value: object) -> str | None:
    """Render a cell as the text a person would read; None for an empty cell.

    Numbers go through Decimal(str(...)) so 2134850 stays "2134850" and
    0.1 stays "0.1" (no binary-float digits). Converting text to money is the
    transform layer's job.
    """
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(str(value)), "f")
    text = str(value).strip()
    return text or None


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]] | None:
    """Return (0-based row index, {field: column index}) of the header row."""
    for i, row in enumerate(rows):
        columns = {}
        for j, cell in enumerate(row):
            field = HEADER_ALIASES.get(_normalise(cell))
            if field and field not in columns:
                columns[field] = j
        if REQUIRED_FIELDS <= columns.keys():
            return i, columns
    return None


def _find_account(rows: list[tuple]) -> str | None:
    for row in rows:
        for cell in row:
            if isinstance(cell, str) and (match := ACCOUNT_PATTERN.search(cell)):
                return match.group(1)
    return None


class ExcelStatementParser(BaseStatementParser):
    name = "excel_money_out_in"

    def can_parse(self, path: Path) -> bool:
        if path.suffix.lower() not in EXCEL_SUFFIXES:
            return False
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except UNREADABLE:
            return False
        try:
            return any(
                _find_header(list(ws.iter_rows(max_row=HEADER_SCAN_ROWS, values_only=True)))
                for ws in wb.worksheets
            )
        finally:
            wb.close()

    def parse(self, path: Path) -> list[RawTransaction]:
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except UNREADABLE as exc:
            raise ParseError(path, f"cannot open workbook: {exc}") from exc
        try:
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                found = _find_header(rows[:HEADER_SCAN_ROWS])
                if found:
                    return self._parse_sheet(path, rows, *found)
        finally:
            wb.close()
        raise ParseError(path, "no sheet has a Posting Date / Money Out / Money In header")

    def _parse_sheet(
        self, path: Path, rows: list[tuple], header_idx: int, columns: dict[str, int]
    ) -> list[RawTransaction]:
        account_id = _find_account(rows[:header_idx])
        if account_id is None:
            raise ParseError(path, "no 'Account number:' line above the header")

        def get(row: tuple, field: str) -> str | None:
            j = columns.get(field)
            return cell_text(row[j]) if j is not None and j < len(row) else None

        transactions = []
        for i, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            date = get(row, "date")
            if date is None or (get(row, "debit") is None and get(row, "credit") is None):
                continue  # balance lines, blank rows, footer text
            try:
                txn = RawTransaction(
                    date=date,
                    description=get(row, "description") or "",
                    ref_id=get(row, "ref_id"),
                    debit=get(row, "debit"),
                    credit=get(row, "credit"),
                    balance=get(row, "balance"),
                    account_id=account_id,
                    source_file=path.name,
                    row_number=i,  # the spreadsheet row, as Excel numbers it
                )
            except ValidationError as exc:
                raise ParseError(path, f"row {i}: {exc.errors()[0]['msg']}") from exc
            transactions.append(txn)
        return transactions
