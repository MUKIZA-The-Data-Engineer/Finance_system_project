"""Write tool results as a tidy Excel workbook: one sheet per table, readable defaults."""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"


@dataclass
class Sheet:
    name: str
    columns: list[str]
    rows: list[list[object]]
    money_columns: tuple[str, ...] = ()


def _sheet_title(name: str, used: set[str]) -> str:
    title = re.sub(r"[\[\]:*?/\\]", "_", name)[:31] or "Sheet"
    base, n = title, 2
    while title.casefold() in used:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title.casefold())
    return title


def _cell(value: object) -> object:
    if isinstance(value, list | tuple):
        return "; ".join(str(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {v}" for k, v in value.items())
    return value


def write_workbook(path: Path, sheets: list[Sheet]) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for sheet in sheets:
        ws = wb.create_sheet(_sheet_title(sheet.name, used))
        ws.append(sheet.columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in sheet.rows:
            ws.append([_cell(v) for v in row])
        ws.freeze_panes = "A2"
        for i, name in enumerate(sheet.columns, start=1):
            letter = get_column_letter(i)
            values = [name] + [_cell(r[i - 1]) for r in sheet.rows[:500] if i - 1 < len(r)]
            width = max((len(str(v)) for v in values if v is not None), default=8)
            ws.column_dimensions[letter].width = min(max(width + 2, 8), 60)
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                if isinstance(cell.value, dt.date | dt.datetime):
                    cell.number_format = DATE_FORMAT
                elif name in sheet.money_columns and isinstance(cell.value, Decimal | int | float):
                    cell.number_format = MONEY_FORMAT
    wb.save(path)
    return path


def default_output(source: Path, suffix: str, out_dir: Path = Path("output")) -> Path:
    """output/<file stem>_<suffix>.xlsx"""
    return Path(out_dir) / f"{Path(source).stem}_{suffix}.xlsx"
