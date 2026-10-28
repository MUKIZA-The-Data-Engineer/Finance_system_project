"""Read any PDF, Excel or CSV file into tables, without knowing its layout in advance.

    tables = read_tables(Path("statement.pdf")).tables

Order of attempts:

1. A known bank layout. If a registered statement parser recognises the file
   (see reconpipe.extract), its rows are used: those parsers handle wrapped
   lines, page breaks and balance lines exactly.
2. Excel (.xlsx, .xlsm, .xls): every sheet; the header row is found by
   scoring the first rows for financial column names (English and French), so
   title blocks above the header are skipped.
3. CSV (and .txt): the encoding (UTF-8, Windows-1252) and the delimiter
   (, ; tab |) are detected, then the header row as for Excel.
4. PDF: ruled or aligned tables via pdfplumber, merged across pages with the
   repeated header removed. A PDF with no detectable table falls back to
   reading text lines that start with a date and end with amounts.

Values are kept as read (text, numbers, dates). Interpreting them is the
standardiser's job. Every row keeps the row or line it came from.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from reconpipe.extract import UnsupportedFileError, default_registry
from reconpipe.extract.base import ParseError

EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
CSV_SUFFIXES = {".csv", ".txt", ".tsv"}
PDF_SUFFIXES = {".pdf"}
SUPPORTED_SUFFIXES = EXCEL_SUFFIXES | CSV_SUFFIXES | PDF_SUFFIXES

# Words that appear in real statement and ledger header rows (accents removed).
HEADER_KEYWORDS = (
    "date", "amount", "debit", "credit", "balance", "description", "details", "narrative",
    "particulars", "memo", "reference", "ref", "transaction", "withdrawal", "deposit",
    "money", "paid", "payee", "account", "value", "posting", "entry", "counterparty",
    "partner", "customer", "vendor", "supplier", "label", "currency",
    "libelle", "montant", "solde", "operation", "valeur", "compte", "piece", "devise",
)  # fmt: skip
HEADER_SCAN_ROWS = 30
MIN_HEADER_KEYWORDS = 2
MIN_TEXT_RATIO = 0.6

ACCOUNT_PATTERN = re.compile(
    r"(?:account\s*(?:number|no\.?|#)?|a/c\s*(?:no\.?)?|compte\s*(?:n[o°]\.?)?|iban)"
    r"\s*[:#.]?\s*([A-Z0-9][A-Z0-9\-/]{4,})",
    re.IGNORECASE,
)

DATE_AT_START = re.compile(
    r"^(\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2,4}|\d{4}-\d{2}-\d{2}"
    r"|\d{1,2}[ \-][A-Za-z]{3,9}[ \-]\d{4})\s+(.*)$"
)
AMOUNT_TOKEN = re.compile(
    r"^\(?[-+]?(?:RWF|FRW|USD|EUR|GBP|KES|UGX|TZS|\$|€|£)?\s?[-+]?"
    r"\d{1,3}(?:[,\s.]\d{3})*(?:[.,]\d{2})?\)?(?:\s?(?:CR|DR))?$",
    re.IGNORECASE,
)


MONEY_SHAPE = re.compile(r"[.,]\d{2}(?:\D*)$|\d[,\s.]\d{3}")


@dataclass
class Table:
    """One table of rows read from a file, with its header and row lineage."""

    source_file: str
    sheet: str
    header: list[str]
    rows: list[list[object]]
    row_numbers: list[int]
    method: str
    account_id: str | None = None
    notes: list[str] = field(default_factory=list)


@dataclass
class ReadResult:
    source_file: str
    kind: str  # pdf | excel | csv
    tables: list[Table] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


class UnreadableFileError(Exception):
    """The file type is not supported, or nothing table-like was found in it."""


# --------------------------------------------------------------------------- #
# Header detection
# --------------------------------------------------------------------------- #


def normalise(text: object) -> str:
    """'Libellé de l'opération' -> 'libelle de l operation'."""
    if text is None:
        return ""
    text = unicodedata.normalize("NFKD", str(text))
    text = "".join(c for c in text if not unicodedata.combining(c)).casefold()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def _score(row: list[object]) -> tuple[int, float]:
    cells = [c for c in row if c is not None and str(c).strip()]
    if not cells:
        return 0, 0.0
    words = {w for c in cells for w in normalise(c).split()}
    hits = {k for k in HEADER_KEYWORDS if k in words}
    text_ratio = sum(isinstance(c, str) for c in cells) / len(cells)
    return len(hits), text_ratio


def find_header(rows: list[list[object]], scan: int = HEADER_SCAN_ROWS) -> int | None:
    """Index of the row that best looks like a header, or None."""
    best, best_score = None, (0, 0.0)
    for i, row in enumerate(rows[:scan]):
        hits, ratio = _score(row)
        if hits >= MIN_HEADER_KEYWORDS and ratio >= MIN_TEXT_RATIO and (hits, ratio) > best_score:
            best, best_score = i, (hits, ratio)
    return best


def _header_names(row: list[object]) -> list[str]:
    names, seen = [], {}
    for i, cell in enumerate(row):
        name = str(cell).strip() if cell is not None and str(cell).strip() else f"column_{i + 1}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return names


def find_account(lines: list[object]) -> str | None:
    for line in lines:
        if line is None:
            continue
        match = ACCOUNT_PATTERN.search(str(line))
        if match:
            return match.group(1).strip().rstrip("-/")
    return None


def _table_from_grid(
    source: str, sheet: str, grid: list[list[object]], numbers: list[int], method: str
) -> Table | None:
    header_idx = find_header(grid)
    if header_idx is None:
        return None
    header = _header_names(grid[header_idx])
    width = len(header)
    rows, row_numbers = [], []
    for values, number in zip(grid[header_idx + 1 :], numbers[header_idx + 1 :], strict=True):
        values = (list(values) + [None] * width)[:width]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        rows.append(values)
        row_numbers.append(number)
    above = [c for row in grid[:header_idx] for c in row]
    return Table(source, sheet, header, rows, row_numbers, method, find_account(above))


# --------------------------------------------------------------------------- #
# Readers
# --------------------------------------------------------------------------- #


def _from_bank_parser(path: Path) -> ReadResult | None:
    registry = default_registry()
    try:
        parser = registry.parser_for(path)
    except UnsupportedFileError:
        return None
    try:
        raw = parser.parse(path)
    except ParseError:
        return None
    header = ["Date", "Description", "Reference", "Debit", "Credit", "Amount", "Balance"]
    rows = [[r.date, r.description, r.ref_id, r.debit, r.credit, r.amount, r.balance] for r in raw]
    account = raw[0].account_id if raw else None
    table = Table(
        path.name,
        parser.name,
        header,
        rows,
        [r.row_number for r in raw],
        f"bank_parser:{parser.name}",
        account,
    )
    kind = "pdf" if path.suffix.lower() in PDF_SUFFIXES else "excel"
    return ReadResult(path.name, kind, [table], [f"known bank layout ({parser.name})"])


def _xls_cell(book, cell) -> object:
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, book.datemode)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        value = cell.value
        return int(value) if float(value).is_integer() else Decimal(str(value))
    return cell.value


def _excel_sheets(path: Path) -> list[tuple[str, list[list[object]]]]:
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        return [
            (s.name, [[_xls_cell(book, c) for c in s.row(r)] for r in range(s.nrows)])
            for s in book.sheets()
        ]
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb]
    finally:
        wb.close()


def _read_excel(path: Path) -> ReadResult:
    result = ReadResult(path.name, "excel")
    try:
        sheets = _excel_sheets(path)
    except Exception as exc:  # openpyxl/xlrd raise many types for a bad file
        raise UnreadableFileError(f"{path.name}: cannot open workbook ({exc})") from exc
    for name, grid in sheets:
        grid = [[v if not isinstance(v, float) else Decimal(str(v)) for v in row] for row in grid]
        table = _table_from_grid(
            path.name, name, grid, list(range(1, len(grid) + 1)), "excel_header"
        )
        if table is None:
            result.notes.append(f"sheet '{name}': no header row found, skipped")
        elif table.rows:
            result.tables.append(table)
    return result


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnreadableFileError("cannot decode text")  # pragma: no cover - latin-1 never fails


def _read_csv(path: Path) -> ReadResult:
    text = _decode(path.read_bytes())
    sample = text[:20000]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    grid, numbers = [], []
    for row in reader:
        grid.append([c if c.strip() else None for c in row])
        numbers.append(reader.line_num)
    table = _table_from_grid(path.name, "csv", grid, numbers, f"csv (delimiter {delimiter!r})")
    if table is None:
        raise UnreadableFileError(f"{path.name}: no header row found in the first lines")
    return ReadResult(path.name, "csv", [table])


def _read_pdf(path: Path) -> ReadResult:
    import pdfplumber
    from pdfplumber.utils.exceptions import PdfminerException

    result = ReadResult(path.name, "pdf")
    try:
        with pdfplumber.open(path) as pdf:
            pages = [(p.extract_tables(), p.extract_text() or "") for p in pdf.pages]
    except (OSError, PdfminerException, ValueError) as exc:
        raise UnreadableFileError(f"{path.name}: cannot open PDF ({exc})") from exc
    all_text = "\n".join(text for _, text in pages)
    account = find_account(all_text.splitlines()[:40])

    merged: Table | None = None
    for tables, _ in pages:
        for grid in tables:
            grid = [[c.replace("\n", " ") if isinstance(c, str) else c for c in r] for r in grid]
            if merged is None:
                merged = _table_from_grid(
                    path.name, "pdf tables", grid, list(range(1, len(grid) + 1)), "pdf_tables"
                )
                continue
            header_idx = find_header(grid)
            same = header_idx is not None and _header_names(grid[header_idx]) == merged.header
            body = grid[header_idx + 1 :] if same else grid
            for values in body:
                if any(v not in (None, "") for v in values):
                    merged.rows.append(
                        (list(values) + [None] * len(merged.header))[: len(merged.header)]
                    )
                    merged.row_numbers.append(len(merged.rows))
    if merged is not None and merged.rows:
        merged.row_numbers = list(range(1, len(merged.rows) + 1))  # printed order
        merged.account_id = account
        result.tables.append(merged)
        return result

    lines = _text_lines(all_text)
    if not lines:
        raise UnreadableFileError(f"{path.name}: no table and no dated transaction lines found")
    widest = max(len(r) for r in lines)
    header = ["Date", "Description", "Amount", "Balance", "Extra"][:widest]
    rows = [(r + [None] * widest)[:widest] for r in lines]
    table = Table(
        path.name, "pdf text", header, rows, list(range(1, len(rows) + 1)), "pdf_text_lines",
        account, ["read from text lines: check the amount signs"],
    )  # fmt: skip
    result.tables.append(table)
    return result


def looks_like_money(token: str) -> bool:
    """'1,250.50', '-300.00', '5 000,00 CR' yes; '7' or 'INV-1151' no (invoice numbers, days)."""
    return bool(AMOUNT_TOKEN.match(token)) and bool(MONEY_SHAPE.search(token))


def _text_lines(text: str) -> list[list[object]]:
    """Lines that start with a date and end with amounts; following text lines continue them."""
    rows: list[list[object]] = []
    for line in (ln.strip() for ln in text.splitlines()):
        match = DATE_AT_START.match(line)
        if not match:
            if rows and line and not looks_like_money(line.split()[-1]) and len(line) < 80:
                rows.append([None, line])  # continuation of the line above
            continue
        date, rest = match.groups()
        tokens = rest.split()
        amounts: list[str] = []
        while tokens:
            candidate = tokens[-1]
            if candidate.upper() in ("CR", "DR") and len(tokens) > 1:
                candidate = f"{tokens[-2]} {candidate}"
                if looks_like_money(candidate):
                    amounts.insert(0, candidate)
                    tokens = tokens[:-2]
                    continue
            if looks_like_money(candidate):
                amounts.insert(0, candidate)
                tokens.pop()
            else:
                break
        if not amounts:
            continue
        rows.append([date, " ".join(tokens), *amounts[-3:]])
    return rows


def read_tables(path: Path) -> ReadResult:
    """Read every table in a PDF, Excel or CSV file."""
    path = Path(path)
    if not path.exists():
        raise UnreadableFileError(f"{path}: file not found")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnreadableFileError(
            f"{path.name}: unsupported file type '{suffix}' (use PDF, Excel or CSV)"
        )
    known = _from_bank_parser(path) if suffix in PDF_SUFFIXES | EXCEL_SUFFIXES else None
    if known is not None:
        return known
    if suffix in EXCEL_SUFFIXES:
        result = _read_excel(path)
    elif suffix in CSV_SUFFIXES:
        result = _read_csv(path)
    else:
        result = _read_pdf(path)
    if not result.tables:
        raise UnreadableFileError(f"{path.name}: no table with a recognisable header found")
    return result


def cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)
