"""Builders for 'unknown' files the toolkit has never seen: no registered parser claims them.

Each builder writes into a directory and returns the path. Values are chosen so
the tests know exactly what a correct read looks like.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas

# date, description, reference, debit, credit, balance (balance follows from 100,000.00)
LINES = [
    ("05/01/2026", "Paiement fournisseur ACME SARL", "VIR001", "1 250,50", "", "98 749,50"),
    ("06/01/2026", "Encaissement client BETA LTD", "VIR002", "", "5 000,00", "103 749,50"),
    ("15/01/2026", "Frais bancaires", "", "2 500,00", "", "101 249,50"),
    ("20/01/2026", "Paiement loyer janvier", "VIR003", "40 000,00", "", "61 249,50"),
]


def french_csv(directory: Path) -> Path:
    """Semicolon CSV, Windows-1252, French headers, '1 250,50' amounts, a title block."""
    lines = [
        "Relevé de compte;;;;;",
        "Compte n°: 123-456-789;;;;;",
        ";;;;;",
        "Date opération;Libellé;Référence;Débit;Crédit;Solde",
        *[";".join(line) for line in LINES],
        "Total;;;43 750,50;5 000,00;",
    ]
    path = directory / "releve_janvier.csv"
    path.write_bytes(("\r\n".join(lines) + "\r\n").encode("cp1252"))
    return path


def english_csv(directory: Path, name: str = "books.csv") -> Path:
    """Comma CSV of a company ledger with a signed amount column."""
    rows = [
        "Entry ID,Date,Reference,Description,Partner,Amount",
        "JE-1,2026-01-04,VIR001,Payment ACME SARL supplies,ACME SARL,-1250.50",
        "JE-2,2026-01-06,VIR002,Receipt BETA LTD invoice 7,BETA LTD,5000.00",
        "JE-3,2026-01-15,,Bank charges,,-2500.00",
        "JE-4,2026-01-21,,Office rent January,Landlord Co,-40000.00",
    ]
    path = directory / name
    path.write_text("\n".join(rows) + "\n", encoding="utf-8")
    return path


def ruled_pdf(directory: Path) -> Path:
    """A ruled table over two pages, header repeated on page two, French columns."""
    path = directory / "banque_x.pdf"
    c = Canvas(str(path), pagesize=A4)
    xs = [40, 110, 300, 370, 440, 500, 570]
    header = ["Date", "Libellé", "Référence", "Débit", "Crédit", "Solde"]
    pages = [LINES[:2], LINES[2:]]
    for page in pages:
        c.drawString(40, 800, "BANQUE X - Relevé")
        c.drawString(40, 785, "Account number: BX-0099-12")
        top, h = 740, 16
        lines = [header, *page]
        for i in range(len(lines) + 1):
            c.line(xs[0], top - i * h, xs[-1], top - i * h)
        for x in xs:
            c.line(x, top, x, top - len(lines) * h)
        for i, cells in enumerate(lines):
            for x, text in zip(xs, cells, strict=False):
                c.drawString(x + 3, top - i * h - 11, text)
        c.showPage()
    c.save()
    return path


def text_pdf(directory: Path) -> Path:
    """No table structure at all: one transaction per text line, amount then balance."""
    path = directory / "plain_statement.pdf"
    c = Canvas(str(path), pagesize=A4)
    c.drawString(40, 800, "Statement of account - A/C No. PL-778899")
    lines = [
        "05/01/2026  Payment ACME SARL  -1,250.50  98,749.50",
        "06/01/2026  Receipt BETA LTD  5,000.00  103,749.50",
        "    invoice 7",
        "15/01/2026  Bank charges  -2,500.00  101,249.50",
    ]
    for i, line in enumerate(lines):
        c.drawString(40, 760 - i * 14, line)
    c.save()
    return path


def messy_xlsx(directory: Path) -> Path:
    """A cover sheet with no table, then a sheet with a title block above the header."""
    wb = Workbook()
    wb.active.title = "Cover"
    wb.active.append(["Monthly export"])
    ws = wb.create_sheet("Transactions")
    ws.append(["ACME Bank"])
    ws.append(["Account No: AC-5555-01"])
    ws.append([])
    ws.append(["Txn Date", "Details", "Withdrawals", "Deposits", "Running Balance"])
    ws.append([datetime(2026, 1, 5), "Payment ACME SARL", 1250.5, None, 98749.5])
    ws.append([datetime(2026, 1, 6), "Receipt BETA LTD", None, 5000, 103749.5])
    ws.append([None, None, None, None, None])
    ws.append([datetime(2026, 1, 15), "Bank charges", 2500, None, 101249.5])
    path = directory / "messy_export.xlsx"
    wb.save(path)
    return path


def old_xls(directory: Path) -> Path:
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("Sheet1")
    date_style = xlwt.easyxf(num_format_str="yyyy-mm-dd")
    header = ["Value Date", "Narrative", "Debit", "Credit", "Balance"]
    rows = [
        (datetime(2026, 1, 5), "Payment ACME SARL", 1250.5, None, 98749.5),
        (datetime(2026, 1, 6), "Receipt BETA LTD", None, 5000, 103749.5),
    ]
    for j, name in enumerate(header):
        sheet.write(0, j, name)
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            if value is None:
                continue
            sheet.write(
                i, j, value, date_style if isinstance(value, datetime) else xlwt.Style.default_style
            )
    path = directory / "old_format.xls"
    book.save(str(path))
    return path
