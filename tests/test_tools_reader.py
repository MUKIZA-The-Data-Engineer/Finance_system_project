from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

import tool_fixtures as fx
from reconpipe.cli import main
from reconpipe.tools.reader import UnreadableFileError, find_header, normalise, read_tables

RAW = Path(__file__).resolve().parent.parent / "data" / "raw"


def test_known_bank_layouts_use_their_own_parser():
    for name, rows in [
        ("bank_a_statement_2026Q1.pdf", 111),
        ("bank_b_statement_2026Q1.pdf", 96),
        ("bank_c_statement_2026Q1.xlsx", 85),
    ]:
        (table,) = read_tables(RAW / name).tables
        assert table.method.startswith("bank_parser:")
        assert len(table.rows) == rows and table.account_id


def test_ledger_header_is_found_generically():
    (table,) = read_tables(RAW / "ledger.xlsx").tables
    assert table.method == "excel_header"
    assert table.header[:3] == ["entry_id", "posting_date", "bank_account"]
    assert len(table.rows) == 296 and table.row_numbers[0] == 2


def test_french_semicolon_csv(tmp_path):
    (table,) = read_tables(fx.french_csv(tmp_path)).tables
    assert table.header == ["Date opération", "Libellé", "Référence", "Débit", "Crédit", "Solde"]
    assert table.account_id == "123-456-789"
    assert table.rows[0][:2] == ["05/01/2026", "Paiement fournisseur ACME SARL"]
    assert table.row_numbers[0] == 5  # the line in the file
    assert table.rows[-1][0] == "Total"  # kept as read; the standardiser drops it


def test_messy_workbook_skips_cover_and_title_block(tmp_path):
    result = read_tables(fx.messy_xlsx(tmp_path))
    (table,) = result.tables
    assert table.sheet == "Transactions" and table.header[0] == "Txn Date"
    assert table.account_id == "AC-5555-01"
    assert len(table.rows) == 3  # the blank row is dropped
    assert table.rows[0][2] == Decimal("1250.5")  # floats become Decimal, never float
    assert any("Cover" in n for n in result.notes)


def test_old_xls(tmp_path):
    (table,) = read_tables(fx.old_xls(tmp_path)).tables
    assert table.header[0] == "Value Date" and len(table.rows) == 2
    assert table.rows[0][0].year == 2026


def test_unknown_ruled_pdf_tables_merge_across_pages(tmp_path):
    (table,) = read_tables(fx.ruled_pdf(tmp_path)).tables
    assert table.method == "pdf_tables"
    assert table.header == ["Date", "Libellé", "Référence", "Débit", "Crédit", "Solde"]
    assert [r[0] for r in table.rows] == [line[0] for line in fx.LINES]  # header not repeated
    assert table.account_id == "BX-0099-12"


def test_pdf_without_a_table_is_read_from_text_lines(tmp_path):
    (table,) = read_tables(fx.text_pdf(tmp_path)).tables
    assert table.method == "pdf_text_lines"
    assert table.rows[0] == ["05/01/2026", "Payment ACME SARL", "-1,250.50", "98,749.50"]
    assert table.rows[2][:2] == [
        None,
        "invoice 7",
    ]  # a continuation line, kept for the standardiser
    assert table.account_id == "PL-778899"


def test_unsupported_and_missing_files(tmp_path):
    (tmp_path / "x.docx").write_text("hi")
    with pytest.raises(UnreadableFileError, match="unsupported"):
        read_tables(tmp_path / "x.docx")
    with pytest.raises(UnreadableFileError, match="not found"):
        read_tables(tmp_path / "missing.csv")
    (tmp_path / "noheader.csv").write_text("1,2,3\n4,5,6\n")
    with pytest.raises(UnreadableFileError, match="no header"):
        read_tables(tmp_path / "noheader.csv")


def test_header_detection_ignores_title_rows():
    grid = [
        ["Statement"],
        ["Period: Jan"],
        ["Date", "Description", "Amount"],
        ["2026-01-01", "x", 5],
    ]
    assert find_header(grid) == 2
    assert normalise("Libellé de l'opération") == "libelle de l operation"


def test_cli_pdf_to_excel(tmp_path, capsys):
    out = tmp_path / "out.xlsx"
    assert main(["pdf-to-excel", str(fx.ruled_pdf(tmp_path)), "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["pdf tables", "about"]
    assert wb["pdf tables"].max_row == 1 + len(fx.LINES)
    assert "Read 4 rows" in capsys.readouterr().out


def test_cli_reports_unreadable_files(tmp_path, capsys):
    (tmp_path / "a.csv").write_text("nothing here\n")
    assert main(["pdf-to-excel", str(tmp_path / "a.csv")]) == 2
    assert "needs a .pdf" in capsys.readouterr().err
