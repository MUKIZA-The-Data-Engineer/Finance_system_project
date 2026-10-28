import datetime as dt
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

import tool_fixtures as fx
import truth
from reconpipe.cli import main
from reconpipe.tools.columns import MappingInputError, map_columns, parse_overrides
from reconpipe.tools.standardizer import (
    MappingError,
    parse_amount,
    standardize_file,
    write_standardized,
)
from reconpipe.transform.standardize import StandardizeError

RAW = truth.ROOT / "data" / "raw"
TODAY = dt.date(2026, 10, 26)
D = Decimal


def amounts(result):
    return [r.amount for r in result.rows]


# --------------------------------------------------------------------------- #
# Column mapping
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    ("header", "expected"),
    [
        (
            ["Date opération", "Libellé", "Référence", "Débit", "Crédit", "Solde"],
            {"date", "description", "reference", "debit", "credit", "balance"},
        ),
        (
            ["Txn Date", "Details", "Withdrawals", "Deposits", "Running Balance"],
            {"date", "description", "debit", "credit", "balance"},
        ),
        (
            ["Posting Date", "Narrative", "Txn Ref", "Money Out", "Money In", "Balance"],
            {"date", "description", "reference", "debit", "credit", "balance"},
        ),
        (
            ["Entry ID", "Date", "Reference", "Description", "Partner", "Amount"],
            {"record_id", "date", "reference", "description", "counterparty", "amount"},
        ),
    ],
)
def test_headers_map_to_standard_fields(header, expected):
    assert set(map_columns(header).by_field) == expected


def test_specific_fields_win_over_general_ones():
    fields = map_columns(["Value Date", "Debit Amount", "Credit Amount", "Amount"]).by_field
    assert fields == {
        "date": "Value Date",
        "debit": "Debit Amount",
        "credit": "Credit Amount",
        "amount": "Amount",
    }


def test_overrides_win_and_bad_overrides_are_rejected():
    mapping = map_columns(
        ["When", "What", "How much"], parse_overrides(["When=date", "How much=amount"])
    )
    assert mapping.by_field == {"date": "When", "amount": "How much"}
    assert mapping.missing == []
    with pytest.raises(MappingInputError, match="unknown field"):
        parse_overrides(["When=day"])
    with pytest.raises(MappingInputError, match="expects"):
        parse_overrides(["date"])


def test_missing_columns_are_reported(tmp_path):
    path = tmp_path / "x.csv"
    path.write_text("Description,Reference,Notes\nrent,R1,x\n")
    with pytest.raises(MappingError, match="could not find date"):
        standardize_file(path)


# --------------------------------------------------------------------------- #
# Values
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        ("1 250,50", "1250.50"),
        ("1.250,50", "1250.50"),
        ("1,250.50", "1250.50"),
        ("€1,234.56", "1234.56"),
        ("$ -99.90", "-99.90"),
        ("(500.00)", "-500.00"),
        ("5 000,00 CR", "5000.00"),
        ("700,00 DR", "-700.00"),
        ("RWF -702,650.00", "-702650.00"),
        (D("12.5"), "12.50"),
        (3, "3.00"),
    ],
)
def test_amounts(value, expected):
    assert parse_amount(value) == D(expected)


def test_blank_and_bad_amounts():
    assert parse_amount(None) is None and parse_amount("  ") is None and parse_amount("-") is None
    with pytest.raises(StandardizeError):
        parse_amount("twelve")


# --------------------------------------------------------------------------- #
# Whole files the toolkit has never seen
# --------------------------------------------------------------------------- #


def test_french_csv(tmp_path):
    r = standardize_file(fx.french_csv(tmp_path), today=TODAY)
    assert amounts(r) == [D("-1250.50"), D("5000.00"), D("-2500.00"), D("-40000.00")]
    assert [x.date for x in r.rows][:2] == [dt.date(2026, 1, 5), dt.date(2026, 1, 6)]
    assert {x.account_id for x in r.rows} == {"123-456-789"}
    assert r.rows[0].balance == D("98749.50") and r.rows[0].reference == "VIR001"
    assert r.summary["summary_lines_skipped"] == 1  # the Total line
    assert r.summary["errors"] == 0


def test_messy_workbook(tmp_path):
    r = standardize_file(fx.messy_xlsx(tmp_path), today=TODAY)
    assert amounts(r) == [D("-1250.50"), D("5000.00"), D("-2500.00")]
    assert r.rows[0].account_id == "AC-5555-01" and r.rows[0].debit == D("1250.50")


def test_unknown_ruled_pdf(tmp_path):
    r = standardize_file(fx.ruled_pdf(tmp_path), today=TODAY)
    assert amounts(r) == [D("-1250.50"), D("5000.00"), D("-2500.00"), D("-40000.00")]
    assert r.rows[-1].balance == D("61249.50")


def test_text_only_pdf_joins_wrapped_lines(tmp_path):
    r = standardize_file(fx.text_pdf(tmp_path), today=TODAY)
    assert amounts(r) == [D("-1250.50"), D("5000.00"), D("-2500.00")]
    assert r.rows[1].description == "Receipt BETA LTD invoice 7"
    assert r.summary["rows_merged_into_previous"] == 1
    assert r.summary["errors"] == 0  # the running balance adds up


def test_old_xls(tmp_path):
    r = standardize_file(fx.old_xls(tmp_path), today=TODAY)
    assert amounts(r) == [D("-1250.50"), D("5000.00")]


def test_ambiguous_dates_are_rejected_until_you_choose(tmp_path):
    path = tmp_path / "amb.csv"
    path.write_text("Date,Description,Amount\n05/01/2026,a,10.00\n06/02/2026,b,20.00\n")
    r = standardize_file(path, today=TODAY)
    assert r.rows == [] and any("ambiguous" in i.problem for i in r.issues)
    r = standardize_file(path, date_order="month-first", today=TODAY)
    assert [x.date for x in r.rows] == [dt.date(2026, 5, 1), dt.date(2026, 6, 2)]


def test_problem_rows_are_listed_not_dropped(tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text(
        "Date,Description,Amount,Balance\n"
        "2026-01-05,ok,10.00,110.00\n"
        "2026-01-06,no amount,,\n"
        "2026-13-40,bad date,5.00,\n"
        "2026-01-07,typo amount,abc,\n"
        "2027-01-01,future,1.00,\n"
        "2026-01-08,zero,0.00,\n"
    )
    r = standardize_file(path, today=TODAY)
    assert [x.description for x in r.rows] == ["ok"]
    problems = " | ".join(i.problem for i in r.issues if i.severity == "error")
    for expected in ("missing amount", "does not match", "not an amount", "date_not_in_future",
                     "amount_non_zero"):  # fmt: skip
        assert expected in problems
    assert any("text where amounts are expected" in i.problem for i in r.issues)


def test_reprints_removed_double_postings_flagged(tmp_path):
    path = tmp_path / "d.csv"
    path.write_text(
        "ID,Date,Description,Amount\n"
        "1,2026-01-05,rent,-100.00\n"
        "2,2026-01-05,rent,-100.00\n"  # same line, different ID: possible double posting
        "2,2026-01-05,rent,-100.00\n"  # exact reprint of the line above
    )
    r = standardize_file(path, today=TODAY)
    assert [x.record_id for x in r.rows] == ["1", "2"]
    assert r.summary["duplicates_removed"] == 1
    assert any("possible double posting" in i.problem for i in r.issues)


# --------------------------------------------------------------------------- #
# The toolkit agrees with the pipeline and with the ground truth
# --------------------------------------------------------------------------- #


def test_generic_path_quarantines_exactly_the_planted_defects():
    rows, quarantined = 0, set()
    for name in ("bank_a_statement_2026Q1.pdf", "bank_b_statement_2026Q1.pdf",
                 "bank_c_statement_2026Q1.xlsx"):  # fmt: skip
        r = standardize_file(RAW / name, today=TODAY)
        rows += len(r.rows)
        quarantined |= {
            (i.source_file, i.row, i.problem.split(":")[0])
            for i in r.issues
            if i.severity == "error"
        }
    assert rows == 288
    assert quarantined == {(d["file"], d["row_number"], d["rule"]) for d in truth.defects()}


def test_ledger_double_postings_are_flagged_with_ids():
    r = standardize_file(RAW / "ledger.xlsx", today=TODAY)
    assert len(r.rows) == 296
    flagged = {i.value for i in r.issues if "double posting" in i.problem}
    assert flagged == {e["entry_id"] for e in truth.EXPECTED["ledger_only"] if e["duplicate_of"]}


def test_standardizing_the_output_again_changes_nothing(tmp_path):
    first = standardize_file(fx.french_csv(tmp_path), today=TODAY)
    out = write_standardized(first, tmp_path / "clean.xlsx")
    second = standardize_file(out, today=TODAY)
    keep = ("date", "description", "reference", "amount", "balance", "account_id")
    assert [[getattr(r, k) for k in keep] for r in second.rows] == [
        [getattr(r, k) for k in keep] for r in first.rows
    ]


# --------------------------------------------------------------------------- #
# Command line
# --------------------------------------------------------------------------- #


def test_cli_standardize(tmp_path, capsys):
    out = tmp_path / "o.xlsx"
    assert main(["standardize", str(fx.french_csv(tmp_path)), "-o", str(out)]) == 0
    wb = load_workbook(out)
    assert wb.sheetnames == ["clean", "issues", "column mapping", "summary"]
    assert wb["clean"].max_row == 5
    assert "4 clean rows" in capsys.readouterr().out


def test_cli_map_and_date_order(tmp_path):
    path = tmp_path / "odd.csv"
    path.write_text("When,What,How much\n05/01/2026,a,10.00\n06/02/2026,b,20.00\n")
    assert main(["standardize", str(path)]) == 2  # no date column it recognises
    args = ["standardize", str(path), "--map", "When=date", "--map", "How much=amount",
            "--date-order", "day-first", "-o", str(tmp_path / "o.xlsx")]  # fmt: skip
    assert main(args) == 0
    ws = load_workbook(tmp_path / "o.xlsx")["clean"]
    assert ws["A2"].value.date() == dt.date(2026, 1, 5)


def test_cli_pdf_to_excel_writes_clean_and_raw(tmp_path):
    out = tmp_path / "p.xlsx"
    assert main(["pdf-to-excel", str(fx.text_pdf(tmp_path)), "-o", str(out)]) == 0
    names = load_workbook(out).sheetnames
    assert names[:4] == ["clean", "issues", "column mapping", "summary"]
    assert names[-1] == "about" and any(n.startswith("extracted") for n in names)
    assert Path(out).exists()
