"""The generator's ground truth must be true: these tests re-derive it from the files."""

import importlib.util
import json
import sys
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path

import pdfplumber
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
COMMITTED = ROOT / "tests" / "expected_results.json"

spec = importlib.util.spec_from_file_location(
    "generate_statements", ROOT / "scripts" / "generate_statements.py"
)
gen = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = gen
spec.loader.exec_module(gen)


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    out = tmp_path_factory.mktemp("raw")
    expected = gen.generate(42, out, out / "expected_results.json")
    return out, expected


@pytest.fixture(scope="module")
def ledger(run):
    out, _ = run
    ws = load_workbook(out / "ledger.xlsx").active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    return {r[0]: dict(zip(header, r, strict=True)) for r in rows[1:]}


def without_paths(expected):
    copy = json.loads(json.dumps(expected))
    for s in copy["statements"]:
        s.pop("file")
    copy["ledger"].pop("file")
    return copy


def test_seed_is_deterministic_and_committed_ground_truth_is_current(run):
    _, expected = run
    committed = json.loads(COMMITTED.read_text(encoding="utf-8"))
    assert without_paths(expected) == without_paths(committed), (
        "tests/expected_results.json is stale: run python scripts/generate_statements.py"
    )


def test_every_discrepancy_type_is_planted(run):
    _, expected = run
    assert all(n >= 1 for n in expected["planted"].values())
    matched = expected["expected_matches"]
    assert (
        matched["tier1_exact"] + matched["tier2_amount_date"] + matched["tier3_fuzzy"]
        == matched["total_matched_bank_lines"]
    )
    bank_lines = expected["totals"]["bank_transactions"]
    planted = expected["planted"]
    assert matched["total_matched_bank_lines"] == (
        bank_lines - planted["amount_typo"] - planted["missing_in_ledger"]
    )


def test_ledger_entries_agree_with_ground_truth(run, ledger):
    _, expected = run
    window = expected["tier2_window_days"]
    for t in expected["bank_transactions"]:
        bank_amount = Decimal(t["amount"])
        bank_date = date.fromisoformat(t["date"])
        entries = [ledger[i] for i in t["ledger_entry_ids"]]
        status, tier = t["expected_status"], t["expected_tier"]
        if status in ("missing_in_ledger", "quarantined"):  # quarantined lines have no entry
            assert entries == []
            continue
        total = sum(Decimal(str(e["amount"])) for e in entries)
        gap = max(abs((e["posting_date"].date() - bank_date).days) for e in entries)
        if status == "amount_mismatch":
            assert total != bank_amount and entries[0]["reference"] == t["ref_id"]
        elif tier == "tier1_exact":
            assert total == bank_amount
            assert all(e["reference"] == t["ref_id"] for e in entries)
        elif tier == "tier2_amount_date":
            assert total == bank_amount and entries[0]["reference"] is None and gap <= window
        else:
            assert tier == "tier3_fuzzy"
            assert total == bank_amount and entries[0]["reference"] is None and gap > window


def test_ledger_only_entries(run, ledger):
    _, expected = run
    statuses = Counter(e["expected_status"] for e in expected["ledger_only"])
    assert statuses["missing_in_bank"] == expected["planted"]["missing_in_bank"]
    assert statuses["duplicate_suspected"] == expected["planted"]["ledger_duplicate"]
    for e in expected["ledger_only"]:
        if e["duplicate_of"]:
            dup, original = ledger[e["entry_id"]], ledger[e["duplicate_of"]]
            assert {**dup, "entry_id": None} == {**original, "entry_id": None}


def test_bank_amounts_are_unique_per_account(run):
    _, expected = run
    seen = Counter(
        (t["account_id"], abs(Decimal(t["amount"]))) for t in expected["bank_transactions"]
    )
    assert max(seen.values()) == 1


def test_pdfs_print_every_reference_and_one_duplicate_row(run):
    out, expected = run
    for s in expected["statements"]:
        if not s["layout"].startswith("pdf"):
            continue
        name = Path(s["file"]).name
        with pdfplumber.open(out / name) as pdf:
            assert len(pdf.pages) == s["pages"] > 1
            text = "\n".join(page.extract_text() for page in pdf.pages)
        rows = [t for t in expected["bank_transactions"] if t["file"] == name]
        for t in rows:
            if t["ref_id"]:
                times = 2 if t["seq"] in s["duplicate_printed_seqs"] else 1
                assert text.count(t["ref_id"]) == times


def test_excel_rows_are_where_ground_truth_says(run):
    out, expected = run
    ws = load_workbook(out / "bank_c_statement_2026Q1.xlsx").active
    rows = [t for t in expected["bank_transactions"] if t["file"].endswith(".xlsx")]
    assert rows
    for t in rows:
        assert ws.cell(row=t["excel_row"], column=1).value == (t["printed_date"] or t["date"])
        assert ws.cell(row=t["excel_row"], column=3).value == t["ref_id"]


def test_one_statement_defect_per_business_rule(run):
    _, expected = run
    defects = expected["quality_defects"]
    assert sorted(d["rule"] for d in defects) == [
        "amount_non_zero",
        "balance_continuity",
        "date_not_in_future",
    ]
    assert expected["expected_quarantine"]["rows"] == len(defects) == 3
    assert len({d["file"] for d in defects}) == 3  # one per bank
    for d in defects:
        if d["seq"] is not None:  # a misprint on a real line with no ledger entry
            (t,) = [
                t
                for t in expected["bank_transactions"]
                if t["file"] == d["file"] and t["seq"] == d["seq"]
            ]
            assert t["expected_status"] == "quarantined" and t["ledger_entry_ids"] == []
