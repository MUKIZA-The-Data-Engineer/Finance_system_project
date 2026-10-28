"""One-step processing and the full workflow, on unseen files and on the real data."""

import datetime as dt

import pytest
from openpyxl import load_workbook

import tool_fixtures as fx
import truth
from reconpipe.cli import main
from reconpipe.tools.odoo import load_config
from reconpipe.tools.reader import UnreadableFileError
from reconpipe.tools.standardizer import MappingError
from reconpipe.tools.workflow import find_files, process_file, run_workflow

RAW = truth.ROOT / "data" / "raw"
TODAY = dt.date(2026, 10, 26)


@pytest.fixture
def cfg(tmp_path):
    path = tmp_path / "odoo.toml"
    path.write_text('[odoo]\njournal = "BNK1"\n')
    return load_config(path)


def test_process_csv_gives_clean_and_odoo(tmp_path, cfg):
    result = process_file(fx.french_csv(tmp_path), tmp_path / "out", cfg, today=TODAY)
    stem = result.source.stem
    assert [f.name for f in result.files] == [f"{stem}_clean.xlsx", f"{stem}_odoo.xlsx"]
    assert result.export.exported_rows == 4
    assert "4 clean rows" in result.describe() and "4 line(s) ready for Odoo" in result.describe()


def test_process_pdf_keeps_the_extracted_rows(tmp_path, cfg):
    result = process_file(fx.ruled_pdf(tmp_path), tmp_path / "out", cfg, today=TODAY)
    extracted, odoo = result.files
    assert extracted.name.endswith("_extracted.xlsx") and odoo.name.endswith("_odoo.xlsx")
    assert any(n.startswith("extracted") for n in load_workbook(extracted).sheetnames)


def test_process_without_odoo(tmp_path):
    result = process_file(fx.messy_xlsx(tmp_path), tmp_path / "out", None, today=TODAY)
    assert result.export is None and len(result.files) == 1


def test_find_files_in_a_folder(tmp_path):
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    fx.french_csv(inbox)
    fx.messy_xlsx(inbox)
    (inbox / "notes.txt").write_text("not a statement")
    (inbox / "~$lock.xlsx").write_text("excel lock file")
    assert len(find_files([inbox])) == 2
    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(UnreadableFileError):
        find_files([empty])


def test_workflow_on_unseen_files(tmp_path, cfg):
    result = run_workflow(
        [fx.french_csv(tmp_path)], [fx.english_csv(tmp_path)], tmp_path / "wf", cfg, today=TODAY
    )
    assert result.reconciliation.summary["match_rate"] == 1.0
    names = sorted(f.relative_to(result.folder).as_posix() for f in result.files)
    assert "reconciliation.xlsx" in names and "summary.txt" in names
    assert any(n.startswith("bank/") and n.endswith("_odoo.xlsx") for n in names)
    assert not any(n.startswith("books/") and n.endswith("_odoo.xlsx") for n in names)
    assert "4 of 4 bank lines matched" in (result.folder / "summary.txt").read_text()


def test_workflow_on_the_real_data_equals_the_pipeline(tmp_path, cfg):
    banks = [RAW / "bank_a_statement_2026Q1.pdf", RAW / "bank_b_statement_2026Q1.pdf",
             RAW / "bank_c_statement_2026Q1.xlsx"]  # fmt: skip
    result = run_workflow(banks, [RAW / "ledger.xlsx"], tmp_path / "wf", cfg, today=TODAY)
    s = result.reconciliation.summary
    assert (s["a_lines"], s["matched_a_lines"]) == (288, 278)
    assert sum(p.export.exported_rows for p in result.bank) == 288  # quarantined lines left out


def test_workflow_stops_on_a_file_it_cannot_map(tmp_path, cfg):
    unreadable = tmp_path / "u.csv"
    unreadable.write_text("When,What,How much\n25/01/2026,a,10.00\n")
    with pytest.raises(MappingError):
        run_workflow([unreadable], [fx.english_csv(tmp_path)], tmp_path / "wf", cfg)


def test_cli_process_folder_and_workflow(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(truth.ROOT)  # the default config/odoo.toml
    inbox = tmp_path / "inbox"
    inbox.mkdir()
    bank, books = fx.french_csv(inbox), fx.english_csv(tmp_path)
    assert main(["process", str(inbox), "--out-dir", str(tmp_path / "out"), "--csv"]) == 0
    printed = capsys.readouterr().out
    assert "Done: 1 of 1 file(s)" in printed and "EXAMPLE codes" in printed
    assert (tmp_path / "out" / bank.stem / f"{bank.stem}_odoo.csv").exists()
    args = ["workflow", "--bank", str(bank), "--books", str(books),
            "--out-dir", str(tmp_path / "wf"), "--no-odoo"]  # fmt: skip
    assert main(args) == 0
    assert "4 of 4 bank lines matched" in capsys.readouterr().out
    assert main(["workflow", "--bank", str(bank)]) == 2
    bad = tmp_path / "bad.txt"
    bad.write_text("x")
    assert main(["process", str(bad), "--out-dir", str(tmp_path / "out")]) == 2
