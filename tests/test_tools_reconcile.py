import datetime as dt
from decimal import Decimal

import pytest
from openpyxl import load_workbook

import tool_fixtures as fx
import truth
from reconpipe.cli import main
from reconpipe.tools.reconcile import (
    Tolerances,
    reconcile_files,
    reconcile_rows,
    write_report,
)
from reconpipe.tools.standardizer import StandardRow

RAW = truth.ROOT / "data" / "raw"
TODAY = dt.date(2026, 10, 26)
BANKS = [RAW / "bank_a_statement_2026Q1.pdf", RAW / "bank_b_statement_2026Q1.pdf",
         RAW / "bank_c_statement_2026Q1.xlsx"]  # fmt: skip
D = Decimal


def row(n, day, amount, desc="x", ref=None, account="ACC", cp=None, rid=None, source="f"):
    return StandardRow(
        date=dt.date(2026, 1, day),
        description=desc,
        amount=D(amount),
        reference=ref,
        account_id=account,
        counterparty=cp,
        record_id=rid,
        source_file=source,
        row_number=n,
    )


def statuses(result):
    return sorted((i.status, i.tier) for i in result.items)


# --------------------------------------------------------------------------- #
# Against the ground truth: the toolkit agrees with the dbt engine item by item
# --------------------------------------------------------------------------- #


@pytest.fixture(scope="module")
def full():
    return reconcile_files(BANKS, [RAW / "ledger.xlsx"], today=TODAY)


def test_every_bank_line_matches_the_ground_truth(full):
    got = {}
    for item in full.items:
        for r in item.a:
            got[(r.source_file, r.row_number)] = (
                item.status,
                item.tier,
                sorted(x.record_id for x in item.b),
            )
    rename = {"missing_in_ledger": "missing_in_b"}
    wrong = []
    for t in truth.EXPECTED["bank_transactions"]:
        if t["expected_status"] == "quarantined":
            continue
        tier = t["expected_tier"] or (
            "tier1_exact" if t["expected_status"] == "amount_mismatch" else None
        )
        want = (
            rename.get(t["expected_status"], t["expected_status"]),
            tier,
            sorted(t["ledger_entry_ids"]),
        )
        if got.get((t["file"], t["row_number"])) != want:
            wrong.append((t["file"], t["seq"], got.get((t["file"], t["row_number"])), want))
    assert wrong == []


def test_ledger_only_items_match_the_ground_truth(full):
    got = {i.b[0].record_id: i.status for i in full.items if not i.a}
    rename = {"missing_in_bank": "missing_in_a"}
    assert got == {
        e["entry_id"]: rename.get(e["expected_status"], e["expected_status"])
        for e in truth.EXPECTED["ledger_only"]
    }


def test_headline_numbers_equal_the_pipeline(full):
    s = full.summary
    assert (s["a_lines"], s["matched_a_lines"], s["match_rate"]) == (288, 278, 0.9653)
    assert s["value_matched"] == D("318125150.00")
    assert s["value_in_exception"] == D("19450550.00")


# --------------------------------------------------------------------------- #
# Tiers and tolerances on small cases
# --------------------------------------------------------------------------- #


def test_files_the_toolkit_has_never_seen(tmp_path):
    result = reconcile_files([fx.french_csv(tmp_path)], [fx.english_csv(tmp_path)], today=TODAY)
    assert statuses(result) == [
        ("matched", "tier1_exact"),
        ("matched", "tier1_exact"),
        ("matched", "tier2_amount_date"),
        ("matched", "tier2_amount_date"),
    ]
    assert result.summary["match_rate"] == 1.0


def test_reference_groups_settle_many_to_one_and_mismatch():
    a = [row(1, 5, "-300.00", ref="PAY-9"), row(2, 6, "-80.00", ref="X1")]
    b = [
        row(1, 5, "-100.00", ref="pay 9"),
        row(2, 5, "-200.00", ref="PAY9"),
        row(3, 6, "-85.00", ref="X1"),
    ]
    result = reconcile_rows(a, b)
    by_status = {i.status: i for i in result.items}
    assert len(by_status["matched"].b) == 2  # references compared without spaces/punctuation
    assert by_status["amount_mismatch"].difference == D("-5.00")


def test_amount_tolerance():
    a, b = [row(1, 5, "-100.00")], [row(1, 5, "-100.40")]
    assert {i.status for i in reconcile_rows(a, b).items} == {"missing_in_a", "missing_in_b"}
    result = reconcile_rows(a, b, Tolerances(amount=D("0.50")))
    (item,) = result.items
    assert item.status == "matched" and "apart" in item.detail and item.confidence < 1


def test_date_tolerance_and_closest_pair_wins():
    a = [row(1, 10, "-50.00")]
    b = [row(1, 13, "-50.00"), row(2, 11, "-50.00")]
    result = reconcile_rows(a, b)
    (matched,) = result.of("matched")
    assert matched.b[0].row_number == 2  # 1 day apart beats 3 days apart
    assert len(result.of("missing_in_a")) == 1
    assert (
        reconcile_rows(a, [row(1, 15, "-50.00")], Tolerances(days=2, description_days=0)).of(
            "matched"
        )
        == []
    )


def test_description_tier_uses_text_and_counterparty():
    a = [row(1, 5, "-900.00", desc="POS 4410 ACME SUPPLIES KIGALI")]
    b = [row(1, 14, "-900.00", desc="Payment ACME Supplies stationery", cp="ACME Supplies")]
    (item,) = reconcile_rows(a, b).items
    assert (item.status, item.tier) == ("matched", "tier3_fuzzy")
    unrelated = [row(1, 14, "-900.00", desc="Office rent", cp="Landlord")]
    assert reconcile_rows(a, unrelated).of("matched") == []


def test_group_amounts_both_directions():
    a = [row(1, 5, "-300.00", desc="batch payment"), row(2, 9, "500.00", desc="a"),
         row(3, 10, "250.00", desc="b")]  # fmt: skip
    b = [row(1, 4, "-100.00", desc="inv 1"), row(2, 6, "-200.00", desc="inv 2"),
         row(3, 9, "750.00", desc="two receipts")]  # fmt: skip
    result = reconcile_rows(a, b)
    groups = sorted((len(i.a), len(i.b)) for i in result.of("matched"))
    assert groups == [(1, 2), (2, 1)]
    assert all(i.tier == "tier4_group_amount" for i in result.of("matched"))


def test_accounts_must_agree_when_both_sides_have_one():
    a, b = [row(1, 5, "-70.00", account="A1")], [row(1, 5, "-70.00", account="B2")]
    assert reconcile_rows(a, b).of("matched") == []
    assert len(reconcile_rows(a, [row(1, 5, "-70.00", account=None)]).of("matched")) == 1


def test_duplicates_on_either_side_are_reported_not_matched():
    a = [row(1, 5, "-70.00", desc="fee"), row(2, 5, "-70.00", desc="fee")]
    b = [row(1, 5, "-70.00", desc="fee", rid="J1"), row(2, 5, "-70.00", desc="fee", rid="J2")]
    result = reconcile_rows(a, b)
    dupes = result.of("duplicate_suspected")
    assert len(dupes) == 2 and {bool(d.a) for d in dupes} == {True, False}
    assert len(result.of("matched")) == 1


# --------------------------------------------------------------------------- #
# Report and command line
# --------------------------------------------------------------------------- #


def test_report_workbook(full, tmp_path):
    wb = load_workbook(write_report(full, tmp_path / "r.xlsx"))
    counts = {n: wb[n].max_row - 1 for n in wb.sheetnames}
    assert counts["Matched"] == 278 and counts["Amount mismatch"] == 6
    assert counts["Missing in books"] == 4 and counts["Missing in bank"] == 3
    assert counts["Duplicate suspected"] == 3 and counts["All items"] == 294
    assert counts["Input issues"] == 7  # 3 quarantined, 1 reprint, 3 double postings flagged


def test_cli_reconcile_two_files(tmp_path, capsys):
    out = tmp_path / "r.xlsx"
    args = [
        "reconcile",
        str(fx.french_csv(tmp_path)),
        str(fx.english_csv(tmp_path)),
        "-o",
        str(out),
    ]
    assert main(args) == 0
    assert "4 of 4 bank lines matched (100.0%)" in capsys.readouterr().out


def test_cli_reconcile_options_and_errors(tmp_path, capsys):
    a, b = fx.french_csv(tmp_path), fx.english_csv(tmp_path)
    args = ["reconcile", "--a", str(a), "--b", str(b), "--a-label", "statement",
            "--b-label", "ledger", "--date-tolerance", "0", "--description-days", "0",
            "-o", str(tmp_path / "r.xlsx")]  # fmt: skip
    assert main(args) == 0
    assert "missing in ledger" in capsys.readouterr().out  # rent is 1 day apart: no longer matched
    assert main(["reconcile", str(a)]) == 2
    assert main(["reconcile", "--a", str(a)]) == 2
