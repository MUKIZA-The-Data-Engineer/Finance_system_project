"""The AI assistant, with a fake Claude client: no network, no API key."""

import json
from types import SimpleNamespace

import anthropic
import pytest
from openpyxl import load_workbook

import tool_fixtures as fx
from reconpipe.cli import main
from reconpipe.tools import ai
from reconpipe.tools.columns import map_columns
from reconpipe.tools.reader import read_tables


class FakeClient:
    """Stands in for anthropic.Anthropic(); records each request."""

    def __init__(self, answer="ok", stop_reason="end_turn", error=None):
        self.requests = []
        self.answer, self.stop_reason, self.error = answer, stop_reason, error
        self.beta = SimpleNamespace(messages=SimpleNamespace(create=self.create))

    def create(self, **kwargs):
        self.requests.append(kwargs)
        if self.error:
            raise self.error
        text = self.answer if isinstance(self.answer, str) else json.dumps(self.answer)
        block = SimpleNamespace(type="text", text=text)
        return SimpleNamespace(content=[block], stop_reason=self.stop_reason)


ODD = "When,What,How much\n25/01/2026,rent,-10.00\n26/01/2026,fees,-20.00\n"
CLAUDE_SAYS = {
    "columns": [
        {"column": "When", "field": "date", "confidence": 0.95},
        {"column": "What", "field": "description", "confidence": 0.9},
        {"column": "How much", "field": "amount", "confidence": 0.9},
        {"column": "Not in the file", "field": "balance", "confidence": 0.9},
    ]
}


@pytest.fixture
def odd(tmp_path):
    path = tmp_path / "odd.csv"
    path.write_text(ODD)
    return path


@pytest.fixture
def fake(monkeypatch):
    client = FakeClient(CLAUDE_SAYS)
    monkeypatch.setattr(ai, "make_client", lambda: client)
    return client


# --------------------------------------------------------------------------- #
# Prompt library
# --------------------------------------------------------------------------- #


def test_library_has_the_prompts_and_every_placeholder_is_known(odd):
    library = ai.prompts()
    assert {"map_columns", "pandas_transform", "odoo_mapping", "review_issues"} <= set(library)
    for name in library:
        text = ai.render(name, odd)
        assert "{" not in text, name  # every placeholder filled in
        assert library[name].title and library[name].description


def test_render_shares_samples_unless_headers_only(odd):
    assert "25/01/2026 | rent" in ai.render("pandas_transform", odd)
    private = ai.render("pandas_transform", odd, headers_only=True)
    assert "rent" not in private and "- How much" in private


def test_review_issues_reports_the_standardiser_findings(tmp_path):
    text = ai.render("review_issues", fx.french_csv(tmp_path))
    assert "-> date" in text and "row " in text


def test_unknown_prompt(odd):
    with pytest.raises(ai.AIError, match="available"):
        ai.render("nope", odd)


# --------------------------------------------------------------------------- #
# Calling Claude
# --------------------------------------------------------------------------- #


def test_request_shape():
    client = FakeClient("answer")
    assert ai.ask_claude("hello", client, schema={"type": "object"}) == "answer"
    (req,) = client.requests
    assert req["model"] == "claude-opus-5-5"
    assert req["fallbacks"] == "default" and req["betas"] == [ai.FALLBACK_BETA]
    assert req["output_config"]["format"]["type"] == "json_schema"
    assert "thinking" not in req  # always on for this model; depth set by effort
    assert req["output_config"]["effort"] == "medium"


@pytest.mark.parametrize(
    ("client", "message"),
    [
        (FakeClient(stop_reason="refusal"), "declined"),
        (FakeClient(stop_reason="max_tokens"), "cut off"),
        (FakeClient(answer=""), "empty"),
        (FakeClient(error=anthropic.APIConnectionError(request=None)), "could not reach"),
    ],
)
def test_failures_become_clear_errors(client, message):
    with pytest.raises(ai.AIError, match=message):
        ai.ask_claude("hello", client)


def test_missing_key_is_a_clear_error(monkeypatch):
    for var in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN"):
        monkeypatch.delenv(var, raising=False)

    def no_key():
        raise anthropic.AnthropicError("no key")

    monkeypatch.setattr(anthropic, "Anthropic", no_key)
    with pytest.raises(ai.AIError, match="ANTHROPIC_API_KEY"):
        ai.make_client()


# --------------------------------------------------------------------------- #
# Column mapping: rules first, Claude for the gaps, remembered afterwards
# --------------------------------------------------------------------------- #


def test_mapper_fills_gaps_and_remembers(odd, fake):
    table = read_tables(odd).tables[0]
    mapper = ai.make_mapper(use_ai=True, client=fake)
    mapping = mapper(table, map_columns(table.header))
    assert mapping.missing == []
    assert {c.source: (c.field, c.how) for c in mapping.choices} == {
        "When": ("date", "ai"),
        "What": ("description", "ai"),
        "How much": ("amount", "ai"),
    }
    # second time: from memory, Claude is not asked again, even without --ai
    again = ai.make_mapper(use_ai=False)(table, map_columns(table.header))
    assert again.missing == [] and len(fake.requests) == 1


def test_rules_and_overrides_keep_priority():
    header = ["Date", "Blurb", "Value"]
    rules = map_columns(header, {"Value": "ignore"})
    merged = ai.merge(rules, {"Date": "balance", "Blurb": "description", "Value": "amount"})
    assert [(c.field, c.how) for c in merged.choices] == [
        ("date", "exact"),
        ("description", "ai"),
        ("ignore", "override"),
    ]


def test_low_confidence_suggestions_are_not_used(odd):
    unsure = {"columns": [{**c, "confidence": 0.2} for c in CLAUDE_SAYS["columns"]]}
    table = read_tables(odd).tables[0]
    mapping = ai.make_mapper(True, FakeClient(unsure))(table, map_columns(table.header))
    assert mapping.missing and not ai.load_memory()


def test_cli_standardize_with_ai(odd, fake, tmp_path, capsys):
    assert main(["standardize", str(odd)]) == 2  # the rules alone cannot read it
    out = tmp_path / "o.xlsx"
    assert main(["standardize", str(odd), "--ai", "--headers-only", "-o", str(out)]) == 0
    assert "Asking Claude" in capsys.readouterr().out
    assert "rent" not in fake.requests[0]["messages"][0]["content"]  # headers only
    assert load_workbook(out)["clean"].max_row == 3
    assert main(["standardize", str(odd), "-o", str(out)]) == 0  # remembered now


def test_cli_ai_commands(odd, fake, tmp_path, capsys):
    assert main(["ai", "list"]) == 0
    assert "pandas_transform" in capsys.readouterr().out
    assert main(["ai", "show", "odoo_mapping", str(odd)]) == 0
    assert "journal_id" in capsys.readouterr().out and not fake.requests
    fake.answer = "Use bank statement lines."
    out = tmp_path / "a.md"
    assert main(["ai", "run", "odoo_mapping", str(odd), "-o", str(out)]) == 0
    assert out.read_text().startswith("Use bank statement lines.")
    fake.answer = CLAUDE_SAYS
    assert main(["ai", "map-columns", str(odd)]) == 0
    assert "remembered" in capsys.readouterr().out
    fake.error = anthropic.APIConnectionError(request=None)
    assert main(["ai", "run", "odoo_mapping", str(odd)]) == 2
