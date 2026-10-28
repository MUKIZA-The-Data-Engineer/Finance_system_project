import csv
import datetime as dt
from collections import defaultdict
from decimal import Decimal

import pytest
from openpyxl import load_workbook

import tool_fixtures as fx
import truth
from reconpipe.cli import main
from reconpipe.tools.odoo import (
    JOURNAL_COLUMNS,
    STATEMENT_COLUMNS,
    OdooConfigError,
    build_export,
    export_file,
    load_config,
    write_export,
    write_template,
)
from reconpipe.tools.standardizer import StandardRow

ROOT = truth.ROOT
D = Decimal


def row(n, amount, desc="Payment ACME", account="ACC-1", cp="ACME SARL", ref="R1"):
    return StandardRow(date=dt.date(2026, 1, n), description=desc, amount=D(amount),
                       reference=ref, account_id=account, counterparty=cp,
                       source_file="s.csv", source_sheet="csv", row_number=n)  # fmt: skip


def write_config(tmp_path, text):
    path = tmp_path / "odoo.toml"
    path.write_text(text, encoding="utf-8")
    return path


# --------------------------------------------------------------------------- #
# Settings
# --------------------------------------------------------------------------- #


def test_committed_config_loads_and_flags_example_codes():
    cfg = load_config(ROOT / "config" / "odoo.toml")
    assert cfg.target == "bank_statement_lines"
    assert cfg.examples_in_use() == ["journal"]


def test_config_journals_per_account_and_overrides(tmp_path):
    path = write_config(tmp_path, '[odoo]\njournal = "BNK1"\n[journals]\n"ACC-2" = "BNK2"\n')
    cfg = load_config(path)
    assert (cfg.journal_for("ACC-1"), cfg.journal_for("ACC-2")) == ("BNK1", "BNK2")
    assert cfg.examples_in_use() == []
    assert load_config(path, journal="OTHER").journal == "OTHER"


def test_bad_config_is_rejected(tmp_path):
    with pytest.raises(OdooConfigError, match="unknown setting"):
        load_config(write_config(tmp_path, "[odoo]\njournl = 'X'\n"))
    with pytest.raises(OdooConfigError, match="target must be"):
        load_config(write_config(tmp_path, "[odoo]\ntarget = 'invoices'\n"))
    with pytest.raises(OdooConfigError):
        load_config(write_config(tmp_path, "[odoo\n"))


# --------------------------------------------------------------------------- #
# The two Odoo imports
# --------------------------------------------------------------------------- #


def test_bank_statement_lines(tmp_path):
    cfg = load_config(write_config(tmp_path, '[odoo]\njournal = "BNK1"\n'))
    export = build_export([row(5, "-1250.50"), row(6, "5000.00", cp=None)], cfg)
    assert export.columns == STATEMENT_COLUMNS
    assert export.lines[0] == [
        "2026-01-05",
        "Payment ACME",
        "ACME SARL",
        "R1",
        D("-1250.50"),
        "BNK1",
    ]
    assert export.exported_rows == 2
    assert any("1 line(s) have no partner" in i.problem for i in export.issues)


def test_journal_entries_balance_line_by_line(tmp_path):
    text = ('[odoo]\ntarget = "journal_entries"\njournal = "BNK1"\nbank_account_code = "512"\n'
            'suspense_account_code = "471"\n[accounts]\n"ACC-2" = "513"\n')  # fmt: skip
    cfg = load_config(write_config(tmp_path, text))
    export = build_export([row(5, "-1250.50"), row(6, "5000.00", account="ACC-2")], cfg)
    assert export.columns == JOURNAL_COLUMNS and len(export.lines) == 4
    entries, current = defaultdict(list), None
    for line in export.lines:
        current = line[0] or current
        entries[current].append(line)
    for lines in entries.values():
        assert sum(x[7] for x in lines) == sum(x[8] for x in lines)  # debits == credits
        assert all(x[1] is None for x in lines[1:])  # only the first line carries the entry
    assert [x[4] for x in export.lines] == ["512", "471", "513", "471"]
    assert export.lines[0][8] == D("1250.50")  # money out: credit the bank account


def test_lines_that_cannot_be_imported_are_left_out(tmp_path):
    cfg = load_config(write_config(tmp_path, '[odoo]\njournal = "BNK1"\npartner_required = true\n'))
    rows = [row(5, "-10.00"), row(6, "-20.00", cp=None), row(7, "0.00"), row(8, "-5.00", desc="")]
    export = build_export(rows, cfg)
    assert export.exported_rows == 1
    problems = sorted(i.problem for i in export.issues if i.severity == "error")
    assert problems == ["amount is zero", "no label (description) for Odoo", "no partner"]


def test_any_file_end_to_end_with_csv(tmp_path):
    cfg = load_config(write_config(tmp_path, '[odoo]\njournal = "BNK1"\n'))
    export = export_file(fx.french_csv(tmp_path), cfg)
    assert export.exported_rows == 4
    xlsx, csv_path = write_export(export, tmp_path / "o.xlsx", csv_too=True)
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["odoo import", "issues", "settings used", "how to import"]
    with csv_path.open(encoding="utf-8") as f:
        rows = list(csv.reader(f))
    assert rows[0] == STATEMENT_COLUMNS and len(rows) == 5


def test_quarantined_statement_lines_never_reach_odoo():
    cfg = load_config(ROOT / "config" / "odoo.toml")
    export = export_file(ROOT / "data" / "raw" / "bank_a_statement_2026Q1.pdf", cfg)
    assert export.exported_rows == 110  # 111 printed, the 0.00 line is quarantined
    assert any(i.problem.startswith("amount_non_zero") for i in export.issues)


# --------------------------------------------------------------------------- #
# Templates and command line
# --------------------------------------------------------------------------- #


@pytest.mark.parametrize(
    ("target", "columns"),
    [("bank_statement_lines", STATEMENT_COLUMNS), ("journal_entries", JOURNAL_COLUMNS)],
)
def test_committed_templates_are_current(tmp_path, target, columns):
    fresh = load_workbook(write_template(target, tmp_path / "t.xlsx"))["odoo import"]
    committed = load_workbook(ROOT / "templates" / f"odoo_{target}_template.xlsx")["odoo import"]
    assert [c.value for c in committed[1]] == [c.value for c in fresh[1]] == columns


def test_cli_odoo_export(tmp_path, capsys, monkeypatch):
    monkeypatch.chdir(ROOT)  # so the default config/odoo.toml is found
    out = tmp_path / "o.xlsx"
    assert main(["odoo-export", str(fx.french_csv(tmp_path)), "-o", str(out), "--csv"]) == 0
    printed = capsys.readouterr().out
    assert "4 line(s) ready for Odoo" in printed and "EXAMPLE codes" in printed
    assert out.with_suffix(".csv").exists()
    assert main(["odoo-export"]) == 2
